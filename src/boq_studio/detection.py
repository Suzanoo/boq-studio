from __future__ import annotations

import re
from dataclasses import replace

from openpyxl.utils import get_column_letter

from .models import SheetConfig
from .normalize import safe_number, safe_text

FIELD_KEYWORDS: dict[str, tuple[str, ...]] = {
    "item_no": ("ลำดับ", "รายการที่", "item", "no.", "no", "sequence", "รหัส"),
    "description": ("รายละเอียด", "รายการ", "description", "work item", "particular"),
    "unit": ("หน่วย", "unit"),
    "qty": ("ปริมาณ", "จำนวน", "qty", "quantity"),
    "material": ("ค่าวัสดุ", "วัสดุ", "material"),
    "labor": ("ค่าแรง", "แรงงาน", "labor", "labour"),
    "amount": ("จำนวนเงิน", "รวมเงิน", "เป็นเงิน", "amount", "total"),
}


def normalized_text(value: object) -> str:
    text = safe_text(value).casefold()
    text = re.sub(r"[\s\n\r\t_/\\:;,.()\[\]-]+", " ", text)
    return text.strip()


def _field_score(text: str, field: str) -> int:
    score = 0
    for keyword in FIELD_KEYWORDS[field]:
        key = keyword.casefold()
        if text == key:
            score = max(score, 8)
        elif key in text:
            score = max(score, 4 if len(key) > 2 else 2)
    return score


def detect_header_row(ws, max_scan_rows: int = 120, max_scan_cols: int = 40) -> int:
    best_row = 1
    best_score = -1
    max_row = min(ws.max_row, max_scan_rows)
    max_col = min(ws.max_column, max_scan_cols)
    for row in range(1, max_row + 1):
        texts = [normalized_text(ws.cell(row, col).value) for col in range(1, max_col + 1)]
        row_score = 0
        matched_fields = 0
        for field in FIELD_KEYWORDS:
            field_best = max((_field_score(text, field) for text in texts), default=0)
            row_score += field_best
            matched_fields += int(field_best > 0)
        nonblank = sum(bool(text) for text in texts)
        row_score += matched_fields * 3 + min(nonblank, 10)
        if matched_fields >= 2 and row_score > best_score:
            best_score = row_score
            best_row = row
    return best_row


def _merged_value(ws, row: int, col: int) -> object:
    value = ws.cell(row, col).value
    if value not in (None, ""):
        return value
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return ws.cell(merged.min_row, merged.min_col).value
    return value


def _header_parts(ws, header_row: int, col: int) -> list[str]:
    # BOQs often use merged two- or three-row headers. Propagate merged labels.
    parts: list[str] = []
    for row in range(max(1, header_row), min(ws.max_row, header_row + 1) + 1):
        text = normalized_text(_merged_value(ws, row, col))
        if text and text not in parts:
            parts.append(text)
    return parts


def _header_text(ws, header_row: int, col: int) -> str:
    return " ".join(_header_parts(ws, header_row, col))


def detect_columns(ws, header_row: int, max_scan_cols: int = 40) -> dict[str, int | list[int] | None]:
    max_col = min(ws.max_column, max_scan_cols)
    headers = {col: _header_text(ws, header_row, col) for col in range(1, max_col + 1)}
    parts = {col: _header_parts(ws, header_row, col) for col in range(1, max_col + 1)}
    detected: dict[str, int | list[int] | None] = {}

    for field in ("item_no", "unit", "qty"):
        ranked = sorted(((_field_score(text, field), -col, col) for col, text in headers.items()), reverse=True)
        score, _, col = ranked[0] if ranked else (0, 0, 0)
        detected[field] = col if score > 0 else None

    # Material and labor should be total-money columns, not unit-rate columns.
    for field in ("material", "labor"):
        ranked: list[tuple[int, int, int]] = []
        for col, text in headers.items():
            score = _field_score(text, field) * 10
            child_text = " ".join(parts[col][1:]) if len(parts[col]) > 1 else ""
            if any(word in child_text for word in ("จำนวนเงิน", "รวมเงิน", "amount", "total")):
                score += 12
            if any(word in child_text for word in ("ราคา หน่วย", "ราคา/หน่วย", "unit price")):
                score -= 5
            ranked.append((score, col, col))  # later column wins ties within a merged group
        score, _, col = max(ranked, default=(0, 0, 0))
        detected[field] = col if score > 0 else None

    # Grand Amount prefers explicit combined/total amount labels.
    ranked_amount: list[tuple[int, int]] = []
    for col, text in headers.items():
        score = _field_score(text, "amount") * 10
        if "รวมค่าวัสดุ" in text and ("แรงงาน" in text or "และแรงงาน" in text):
            score += 80
        if "จำนวนเงิน" in text:
            score += 5
        ranked_amount.append((score, col))
    score, col = max(ranked_amount, default=(0, 0))
    detected["amount"] = col if score > 0 else None

    description_ranked = sorted(
        ((_field_score(text, "description"), -col, col) for col, text in headers.items()),
        reverse=True,
    )
    score, _, primary_desc = description_ranked[0] if description_ranked else (0, 0, 0)
    description_cols: list[int] = [primary_desc] if score > 0 else []

    # Include adjacent text-heavy columns before Qty/Unit; split descriptions are common.
    right_limit_candidates = [col for col in (detected.get("qty"), detected.get("unit")) if isinstance(col, int)]
    right_limit = min(right_limit_candidates, default=max_col + 1)
    if primary_desc:
        for col in range(primary_desc + 1, right_limit):
            text_count = 0
            number_count = 0
            for row in range(header_row + 1, min(ws.max_row, header_row + 80) + 1):
                value = ws.cell(row, col).value
                if safe_number(value) is not None:
                    number_count += 1
                elif safe_text(value):
                    text_count += 1
            if text_count > number_count and text_count >= 2:
                description_cols.append(col)

    if not description_cols:
        candidates: list[tuple[int, int]] = []
        for col in range(1, max_col + 1):
            text_length = 0
            for row in range(header_row + 1, min(ws.max_row, header_row + 40) + 1):
                value = ws.cell(row, col).value
                if safe_text(value) and safe_number(value) is None:
                    text_length += len(safe_text(value))
            candidates.append((text_length, col))
        candidates.sort(reverse=True)
        if candidates and candidates[0][0] > 0:
            description_cols = [candidates[0][1]]
    detected["description_cols"] = sorted(set(description_cols))
    return detected


def detect_first_data_row(ws, header_row: int, columns: dict[str, int | list[int] | None]) -> int:
    description_cols = list(columns.get("description_cols") or [])
    useful_cols = [
        columns.get("item_no"), *description_cols, columns.get("unit"), columns.get("qty"),
        columns.get("material"), columns.get("labor"), columns.get("amount"),
    ]
    useful = sorted({int(col) for col in useful_cols if isinstance(col, int)})
    start = min(ws.max_row, header_row + 1)
    end = min(ws.max_row, header_row + 250)

    # Find the first dense detail row. Summary rows usually have only 2-4 populated fields.
    dense_row: int | None = None
    for row in range(start, end + 1):
        values = [ws.cell(row, col).value for col in useful]
        nonblank = sum(bool(safe_text(value)) for value in values)
        numeric = sum(safe_number(value) is not None for value in values)
        description = " ".join(safe_text(ws.cell(row, col).value) for col in description_cols).strip()
        if description and nonblank >= 5 and numeric >= 2:
            dense_row = row
            break

    if dense_row is not None:
        # Keep nearby parent WBS headings immediately above the detail block.
        candidates: list[int] = []
        for row in range(dense_row - 1, max(header_row, dense_row - 10), -1):
            description = " ".join(safe_text(ws.cell(row, col).value) for col in description_cols).strip()
            item_col = columns.get("item_no")
            item_no = safe_text(ws.cell(row, item_col).value) if isinstance(item_col, int) else ""
            numeric = sum(safe_number(ws.cell(row, col).value) is not None for col in useful)
            if description and item_no and numeric <= 1:
                candidates.append(row)
        return min(candidates) if candidates else dense_row

    for row in range(start, end + 1):
        values = [ws.cell(row, col).value for col in useful]
        nonblank = sum(bool(safe_text(value)) for value in values)
        description = " ".join(safe_text(ws.cell(row, col).value) for col in description_cols).strip()
        if description and nonblank >= 2:
            return row
    return header_row + 1


def auto_detect_config(ws, sheet_name: str, initial: SheetConfig | None = None) -> SheetConfig:
    header_row = detect_header_row(ws)
    columns = detect_columns(ws, header_row)
    start_row = detect_first_data_row(ws, header_row, columns)
    base = initial or SheetConfig(sheet_name=sheet_name)
    return replace(
        base,
        sheet_name=sheet_name,
        header_row=header_row,
        start_row=start_row,
        item_no_col=columns.get("item_no") if isinstance(columns.get("item_no"), int) else None,
        description_cols=list(columns.get("description_cols") or []),
        unit_col=columns.get("unit") if isinstance(columns.get("unit"), int) else None,
        qty_col=columns.get("qty") if isinstance(columns.get("qty"), int) else None,
        material_col=columns.get("material") if isinstance(columns.get("material"), int) else None,
        labor_col=columns.get("labor") if isinstance(columns.get("labor"), int) else None,
        amount_col=columns.get("amount") if isinstance(columns.get("amount"), int) else None,
    )


def mapping_summary(cfg: SheetConfig) -> str:
    def letter(col: int | None) -> str:
        return get_column_letter(col) if col else "—"

    desc = "+".join(get_column_letter(col) for col in cfg.description_cols) or "—"
    return (
        f"Header {cfg.header_row}, Data {cfg.start_row} | "
        f"Desc {desc}, Unit {letter(cfg.unit_col)}, Qty {letter(cfg.qty_col)}, "
        f"Material {letter(cfg.material_col)}, Labor {letter(cfg.labor_col)}, Amount {letter(cfg.amount_col)}"
    )
