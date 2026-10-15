from __future__ import annotations

import copy
import platform
import queue
from dataclasses import replace
import threading
import shutil
import tempfile
from datetime import datetime, timezone
from typing import Callable
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .core.classification import RowType
from .commands import CommandID, CommandRegistry
from .core.review import ClassificationReviewService
from .core.extraction import SheetExtractor
from .core.classification import RowClassifier
from .core.tree import DraftTree, DraftTreeBuilder, TreeNodeType
from .core.tree_edit import DraftTreeEditor, TreeEditError
from .detection import auto_detect_config, mapping_summary
from .models import NormalizeRequest, SheetConfig
from .normalize import OUTPUT_HEADERS, normalize_workbook, normalized_preview_rows, safe_text
from .templates import default_template_path, load_template, save_template, template_directory
from .workflow import SheetState, sheet_state, workflow_status
from .workspace import WorkspaceKind, WorkspaceServices, WorkspaceSpec, WorkspaceWindow
from .shortcuts import is_text_input_widget, shortcuts_for
from .native_os import add_recent_file, is_macos, layout_state_path, load_recent_files, recent_files_path, save_recent_files, view_state_path
from .tree_ui import ACTION_COMMANDS, EDIT_COMMANDS, CommandBarMode, command_bar_mode, editing_grid, summarize_tree
from .layout import GridRegion, LayoutManager, LayoutStateStore, PackRegion
from .view_state import ViewStateManager, ViewStateStore
from .desktop_ui import (
    TREE_CONTEXT_MENU, TREE_EDIT_MENU, TREE_FILE_MENU, TREE_MENU, TREE_TOOLBAR, TREE_VIEW_MENU,
    normalized_shortcut, tooltip_text,
)
from .tooltip import TooltipManager
from .appearance import enable_windows_dpi_awareness, install_appearance
from .session import (
    SESSION_VERSION,
    archive_session,
    config_to_dict,
    tree_to_dict,
    session_token,
    workbook_fingerprint,
    config_from_dict,
    load_session,
    save_session,
    session_path_for,
    tree_from_dict,
)

NONE_LABEL = "— None —"


class SheetMappingDialog(WorkspaceWindow):
    def __init__(self, parent: tk.Misc, input_path: Path, sheet_name: str, initial: SheetConfig | None = None):
        super().__init__(
            parent,
            WorkspaceSpec(
                workspace_id="sheet_mapping",
                kind=WorkspaceKind.MAPPING,
                title=f"Map columns — {sheet_name}",
                geometry="1220x790",
                min_size=(980, 650),
                surfaces=("workspace", "command_bar"),
            ),
        )
        self.result: SheetConfig | None = None
        self.input_path = input_path
        self.sheet_name = sheet_name
        self.initial = initial or SheetConfig(sheet_name=sheet_name)
        self.current_max_col = 0

        self.header_row_var = tk.IntVar(value=self.initial.header_row)
        self.start_row_var = tk.IntVar(value=self.initial.start_row)
        self.column_vars: dict[str, tk.StringVar] = {}
        self.description_list: tk.Listbox
        self.preview_tree: ttk.Treeview
        self.normalized_tree: ttk.Treeview
        self._build()
        self._load_preview(self.initial)

    def _build(self) -> None:
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)

        top = ttk.LabelFrame(root, text="1. Row settings", padding=10)
        top.pack(fill="x")
        ttk.Label(top, text="Header row").grid(row=0, column=0, sticky="w")
        ttk.Spinbox(top, from_=1, to=5000, textvariable=self.header_row_var, width=8).grid(row=0, column=1, padx=(6, 20))
        ttk.Label(top, text="First data row").grid(row=0, column=2, sticky="w")
        ttk.Spinbox(top, from_=1, to=5000, textvariable=self.start_row_var, width=8).grid(row=0, column=3, padx=(6, 20))
        ttk.Button(top, text="Refresh source", command=self._load_preview).grid(row=0, column=4, padx=(0, 8))
        ttk.Button(top, text="Auto Detect", command=self._auto_detect, style="Accent.TButton").grid(row=0, column=5)

        mapping = ttk.LabelFrame(root, text="2. Column mapping", padding=10)
        mapping.pack(fill="x", pady=10)
        fields = [
            ("item_no", "Item No. / Existing WBS"),
            ("unit", "Unit"),
            ("qty", "Qty"),
            ("material", "Material"),
            ("labor", "Labor"),
            ("amount", "Amount"),
        ]
        for idx, (key, label) in enumerate(fields):
            ttk.Label(mapping, text=label).grid(row=0, column=idx, sticky="w", padx=4)
            var = tk.StringVar(value=NONE_LABEL)
            self.column_vars[key] = var
            combo = ttk.Combobox(mapping, textvariable=var, state="readonly", width=18)
            combo.grid(row=1, column=idx, padx=4, pady=(3, 0), sticky="ew")
            setattr(self, f"{key}_combo", combo)
            mapping.columnconfigure(idx, weight=1)

        body = ttk.Panedwindow(root, orient="horizontal")
        body.pack(fill="both", expand=True)
        desc_frame = ttk.LabelFrame(body, text="3. Description columns", padding=8)
        preview_frame = ttk.LabelFrame(body, text="4. Preview", padding=8)
        body.add(desc_frame, weight=1)
        body.add(preview_frame, weight=5)

        self.description_list = tk.Listbox(desc_frame, selectmode="extended", exportselection=False, width=30)
        self.description_list.pack(fill="both", expand=True)
        ttk.Label(desc_frame, text="Select one or more columns. Values are merged with spaces.", wraplength=220).pack(fill="x", pady=(6, 0))

        notebook = ttk.Notebook(preview_frame)
        notebook.pack(fill="both", expand=True)
        source_tab = ttk.Frame(notebook)
        normalized_tab = ttk.Frame(notebook)
        notebook.add(source_tab, text="Source")
        notebook.add(normalized_tab, text="Normalized")

        self.preview_tree = self._create_tree(source_tab)
        self.normalized_tree = self._create_tree(normalized_tab)

        preview_actions = ttk.Frame(preview_frame)
        preview_actions.pack(fill="x", pady=(8, 0))
        ttk.Button(preview_actions, text="Refresh normalized preview", command=self._load_normalized_preview).pack(side="right")

        buttons = ttk.Frame(root)
        buttons.pack(fill="x", pady=(10, 0))
        ttk.Button(buttons, text="Cancel", command=self.destroy).pack(side="right")
        ttk.Button(buttons, text="Save mapping", command=self._save, style="Accent.TButton").pack(side="right", padx=(0, 8))
        ttk.Separator(buttons, orient="vertical").pack(side="right", fill="y", padx=8)
        ttk.Button(buttons, text="Save template…", command=self._save_template).pack(side="right")
        ttk.Button(buttons, text="Load template…", command=self._load_template).pack(side="right", padx=(0, 6))

    @staticmethod
    def _create_tree(parent: ttk.Frame) -> ttk.Treeview:
        tree = ttk.Treeview(parent, show="headings", height=18)
        xbar = ttk.Scrollbar(parent, orient="horizontal", command=tree.xview)
        ybar = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
        tree.configure(xscrollcommand=xbar.set, yscrollcommand=ybar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        parent.rowconfigure(0, weight=1)
        parent.columnconfigure(0, weight=1)
        return tree

    def _read_sheet(self):
        wb = load_workbook(self.input_path, read_only=False, data_only=True)
        return wb, wb[self.sheet_name]

    def _load_preview(self, config: SheetConfig | None = None) -> None:
        try:
            wb, ws = self._read_sheet()
            max_col = min(ws.max_column, 60)
            self.current_max_col = max_col
            header_row = max(1, self.header_row_var.get())
            choices = [NONE_LABEL]
            desc_choices: list[str] = []
            for col in range(1, max_col + 1):
                header = safe_text(ws.cell(header_row, col).value)
                label = f"{get_column_letter(col)} — {header or '(blank)'}"
                choices.append(label)
                desc_choices.append(label)
            for key in self.column_vars:
                getattr(self, f"{key}_combo")["values"] = choices

            previous = config or self._config_from_widgets(require_description=False)
            self.description_list.delete(0, "end")
            for label in desc_choices:
                self.description_list.insert("end", label)
            self._apply_config_to_widgets(previous, choices)

            columns = ["ROW"] + [get_column_letter(c) for c in range(1, max_col + 1)]
            self._configure_tree(self.preview_tree, columns, default_width=82)
            preview_start = max(1, header_row - 2)
            preview_end = min(ws.max_row, preview_start + 44)
            for row in range(preview_start, preview_end + 1):
                values = [row] + [safe_text(ws.cell(row, c).value)[:100] for c in range(1, max_col + 1)]
                self.preview_tree.insert("", "end", values=values)
            wb.close()
            self._load_normalized_preview(show_errors=False)
        except Exception as exc:
            messagebox.showerror("Preview error", str(exc), parent=self)

    @staticmethod
    def _configure_tree(tree: ttk.Treeview, columns: list[str], default_width: int = 110) -> None:
        tree.delete(*tree.get_children())
        tree["columns"] = columns
        for col in columns:
            tree.heading(col, text=col)
            width = 58 if col in {"ROW", "WBS Level"} else default_width
            if col in {"Description", "WBS-1", "WBS-2", "WBS-3", "WBS-4"}:
                width = 180
            tree.column(col, width=width, stretch=False)

    def _auto_detect(self) -> None:
        try:
            wb, ws = self._read_sheet()
            cfg = auto_detect_config(ws, self.sheet_name, self._config_from_widgets(require_description=False))
            wb.close()
            self.header_row_var.set(cfg.header_row)
            self.start_row_var.set(cfg.start_row)
            self._load_preview(cfg)
            messagebox.showinfo("Auto Detect", "Suggested mapping applied. Review the source and normalized previews before saving.", parent=self)
        except Exception as exc:
            messagebox.showerror("Auto Detect failed", str(exc), parent=self)

    def _load_normalized_preview(self, show_errors: bool = True) -> None:
        try:
            cfg = self._config_from_widgets(require_description=True)
            wb, ws = self._read_sheet()
            rows = normalized_preview_rows(ws, cfg, limit=80)
            wb.close()
            self._configure_tree(self.normalized_tree, OUTPUT_HEADERS, default_width=105)
            for row in rows:
                values = ["" if value is None else value for value in row]
                self.normalized_tree.insert("", "end", values=values)
            if not rows and show_errors:
                messagebox.showwarning("No normalized rows", "No output rows were detected. Check First data row and column mapping.", parent=self)
        except ValueError as exc:
            if show_errors:
                messagebox.showwarning("Mapping incomplete", str(exc), parent=self)
        except Exception as exc:
            if show_errors:
                messagebox.showerror("Normalized preview error", str(exc), parent=self)

    @staticmethod
    def _column_from_label(label: str) -> int | None:
        if not label or label.startswith("—"):
            return None
        letters = label.split(" — ", 1)[0]
        value = 0
        for char in letters:
            value = value * 26 + ord(char.upper()) - 64
        return value

    def _config_from_widgets(self, require_description: bool = True) -> SheetConfig:
        description_cols = [index + 1 for index in self.description_list.curselection()]
        if require_description and not description_cols:
            raise ValueError("Select at least one Description column.")
        return SheetConfig(
            sheet_name=self.sheet_name,
            header_row=max(1, self.header_row_var.get()),
            start_row=max(1, self.start_row_var.get()),
            item_no_col=self._column_from_label(self.column_vars["item_no"].get()),
            description_cols=description_cols,
            unit_col=self._column_from_label(self.column_vars["unit"].get()),
            qty_col=self._column_from_label(self.column_vars["qty"].get()),
            material_col=self._column_from_label(self.column_vars["material"].get()),
            labor_col=self._column_from_label(self.column_vars["labor"].get()),
            amount_col=self._column_from_label(self.column_vars["amount"].get()),
            max_wbs_level=self.initial.max_wbs_level,
            classification_overrides=dict(self.initial.classification_overrides),
            wbs_level_overrides=dict(self.initial.wbs_level_overrides),
        )

    def _apply_config_to_widgets(self, cfg: SheetConfig, choices: list[str] | None = None) -> None:
        self.header_row_var.set(cfg.header_row)
        self.start_row_var.set(cfg.start_row)
        choices = choices or list(getattr(self, "item_no_combo")["values"])
        column_map = {
            "item_no": cfg.item_no_col,
            "unit": cfg.unit_col,
            "qty": cfg.qty_col,
            "material": cfg.material_col,
            "labor": cfg.labor_col,
            "amount": cfg.amount_col,
        }
        for key, col in column_map.items():
            self.column_vars[key].set(choices[col] if col and col < len(choices) else NONE_LABEL)
        self.description_list.selection_clear(0, "end")
        for col in cfg.description_cols:
            if 1 <= col <= self.description_list.size():
                self.description_list.selection_set(col - 1)

    def _save_template(self) -> None:
        try:
            cfg = self._config_from_widgets(require_description=True)
        except ValueError as exc:
            messagebox.showwarning("Mapping incomplete", str(exc), parent=self)
            return
        suggested = default_template_path(self.input_path, self.sheet_name)
        filename = filedialog.asksaveasfilename(
            parent=self,
            title="Save mapping template",
            initialdir=suggested.parent,
            initialfile=suggested.name,
            defaultextension=".json",
            filetypes=[("BOQ mapping", "*.json")],
        )
        if filename:
            save_template(Path(filename), cfg)
            messagebox.showinfo("Template saved", filename, parent=self)

    def _load_template(self) -> None:
        filename = filedialog.askopenfilename(
            parent=self,
            title="Load mapping template",
            initialdir=template_directory(),
            filetypes=[("BOQ mapping", "*.json"), ("All files", "*.*")],
        )
        if not filename:
            return
        try:
            cfg = load_template(Path(filename), sheet_name=self.sheet_name)
            self.header_row_var.set(cfg.header_row)
            self.start_row_var.set(cfg.start_row)
            self._load_preview(cfg)
        except Exception as exc:
            messagebox.showerror("Cannot load template", str(exc), parent=self)

    def _save(self) -> None:
        try:
            self.result = self._config_from_widgets(require_description=True)
        except ValueError as exc:
            messagebox.showwarning("Missing mapping", str(exc), parent=self)
            return
        self.destroy()


class ClassificationReviewDialog(WorkspaceWindow):
    EDITABLE_TYPES = (RowType.WBS, RowType.ITEM, RowType.IGNORE)

    def __init__(self, parent: tk.Misc, input_path: Path, config: SheetConfig):
        super().__init__(
            parent,
            WorkspaceSpec(
                workspace_id="classification_review",
                kind=WorkspaceKind.CLASSIFICATION_REVIEW,
                title=f"Review row types — {config.sheet_name}",
                geometry="1320x760",
                min_size=(980, 620),
                surfaces=("workspace", "command_bar", "status_bar"),
            ),
        )
        self.input_path = input_path
        self.config = config
        self.result: tuple[dict[int, str], dict[int, int]] | None = None
        self.rows_by_id: dict[str, object] = {}
        self.overrides = dict(config.classification_overrides)
        self.level_overrides = dict(config.wbs_level_overrides)
        self.filter_var = tk.StringVar(value="Review needed")
        self.status_var = tk.StringVar(value="")
        self.type_status_var = tk.StringVar(value="")
        self.editor_state_var = tk.StringVar(value="Ready")
        self.location_var = tk.StringVar(value=f"Workbook: {self.input_path.name}    |    Sheet: {self.config.sheet_name}")
        self._command_bar_mode: CommandBarMode | None = None
        self.commands = CommandRegistry()
        self._register_commands()
        self._build()
        self._load_rows()

    def _build(self) -> None:
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)

        top = ttk.Frame(root)
        top.pack(fill="x", pady=(0, 8))
        ttk.Label(top, text="Filter").pack(side="left")
        filter_box = ttk.Combobox(
            top,
            textvariable=self.filter_var,
            state="readonly",
            values=("Review needed", "All rows", "Overrides only", "UNKNOWN only"),
            width=18,
        )
        filter_box.pack(side="left", padx=6)
        filter_box.bind("<<ComboboxSelected>>", lambda _event: self._render())
        ttk.Label(
            top,
            text="Select rows, then assign WBS-2 / WBS-3 / WBS-4, ITEM or IGNORE.",
        ).pack(side="left", padx=12)

        columns = ("row", "item", "description", "auto", "final", "level", "confidence", "unit", "qty", "amount", "reason")
        self.tree = ttk.Treeview(root, columns=columns, show="headings", selectmode="extended")
        headings = {
            "row": "Source Row", "item": "Item No.", "description": "Description",
            "auto": "Auto", "final": "Final", "level": "WBS Level", "confidence": "Confidence",
            "unit": "Unit", "qty": "Qty", "amount": "Amount", "reason": "Reason",
        }
        widths = {"row": 85, "item": 110, "description": 330, "auto": 80, "final": 80, "level": 85,
                  "confidence": 90, "unit": 70, "qty": 90, "amount": 120, "reason": 300}
        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(col, width=widths[col], stretch=False, anchor="w")
        ybar = ttk.Scrollbar(root, orient="vertical", command=self.tree.yview)
        xbar = ttk.Scrollbar(root, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        self.tree.pack(fill="both", expand=True, side="left")
        ybar.pack(fill="y", side="right")
        xbar.pack(fill="x", side="bottom")

        actions = ttk.Frame(self, padding=(12, 0, 12, 12))
        actions.pack(fill="x")
        ttk.Button(actions, text="WBS-2", command=lambda: self._set_wbs_level(2)).pack(side="left")
        ttk.Button(actions, text="WBS-3", command=lambda: self._set_wbs_level(3)).pack(side="left", padx=(6, 0))
        ttk.Button(actions, text="WBS-4", command=lambda: self._set_wbs_level(4)).pack(side="left", padx=(6, 0))
        ttk.Button(actions, text="Mark ITEM", command=lambda: self._set_selected(RowType.ITEM)).pack(side="left", padx=6)
        ttk.Button(actions, text="Ignore", command=lambda: self._set_selected(RowType.IGNORE)).pack(side="left")
        ttk.Button(actions, text="Reset Auto", command=self._reset_selected).pack(side="left", padx=6)
        ttk.Label(actions, textvariable=self.status_var).pack(side="left", padx=14)
        ttk.Button(actions, text="Cancel", command=self.destroy).pack(side="right")
        ttk.Button(actions, text="Save review", command=self._save, style="Accent.TButton").pack(side="right", padx=6)

    def _load_rows(self) -> None:
        wb = load_workbook(self.input_path, read_only=False, data_only=True)
        try:
            ws = wb[self.config.sheet_name]
            cfg = replace(
                self.config,
                classification_overrides=dict(self.overrides),
                wbs_level_overrides=dict(self.level_overrides),
            )
            rows = ClassificationReviewService().build_rows(ws, cfg)
        finally:
            wb.close()
        self.all_rows = rows
        self._render()

    def _visible(self, row) -> bool:
        mode = self.filter_var.get()
        if mode == "All rows":
            return True
        if mode == "Overrides only":
            return row.source_row in self.overrides or row.source_row in self.level_overrides
        if mode == "UNKNOWN only":
            return row.auto_type == RowType.UNKNOWN
        return (
            row.auto_type == RowType.UNKNOWN
            or row.confidence < 0.80
            or row.source_row in self.overrides
            or row.source_row in self.level_overrides
        )

    def _render(self) -> None:
        self.tree.delete(*self.tree.get_children())
        self.rows_by_id.clear()
        visible = 0
        for row in self.all_rows:
            if not self._visible(row):
                continue
            final = RowType(self.overrides.get(row.source_row, row.auto_type.value))
            iid = str(row.source_row)
            self.rows_by_id[iid] = row
            level = self.level_overrides.get(row.source_row, row.wbs_level if final == RowType.WBS else None)
            self.tree.insert("", "end", iid=iid, values=(
                row.source_row, row.item_no, row.description, row.auto_type.value, final.value,
                "" if level is None else f"WBS-{level}", f"{row.confidence:.0%}", row.unit, "" if row.qty is None else row.qty,
                "" if row.amount is None else row.amount, row.reason,
            ))
            visible += 1
        override_rows = set(self.overrides) | set(self.level_overrides)
        self.status_var.set(f"{visible} shown · {len(override_rows)} reviewed rows")

    def _set_selected(self, row_type: RowType) -> None:
        for iid in self.tree.selection():
            source_row = int(iid)
            self.overrides[source_row] = row_type.value
            if row_type != RowType.WBS:
                self.level_overrides.pop(source_row, None)
        self._render()

    def _set_wbs_level(self, level: int) -> None:
        for iid in self.tree.selection():
            source_row = int(iid)
            self.overrides[source_row] = RowType.WBS.value
            self.level_overrides[source_row] = level
        self._render()

    def _reset_selected(self) -> None:
        for iid in self.tree.selection():
            source_row = int(iid)
            self.overrides.pop(source_row, None)
            self.level_overrides.pop(source_row, None)
        self._render()

    def _save(self) -> None:
        self.result = (dict(self.overrides), dict(self.level_overrides))
        self.destroy()


class DraftTreeDialog(WorkspaceWindow):
    """Minimal MS-5.4 editor for the automatically detected draft WBS tree."""

    def __init__(self, parent: tk.Misc, input_path: Path, config: SheetConfig, initial_tree: DraftTree | None = None,
                 on_commit: Callable[[DraftTree, str], bool] | None = None):
        layout = LayoutManager(
            "draft_tree",
            state_store=LayoutStateStore(layout_state_path(platform.system())),
        )
        services = WorkspaceServices(
            commands=CommandRegistry(),
            layout=layout,
        )
        super().__init__(
            parent,
            WorkspaceSpec(
                workspace_id="draft_tree",
                kind=WorkspaceKind.DRAFT_TREE,
                title=f"Draft WBS tree — {config.sheet_name}",
                geometry="1050x760",
                min_size=(760, 520),
                surfaces=("menu", "toolbar", "workspace", "sidebar", "command_bar", "status_bar"),
            ),
            services=services,
        )
        self.input_path = input_path
        self.config = config
        self.result: DraftTree | None = None
        self.result_mode: str | None = None
        self.dirty = False
        self.undo_stack: list[DraftTree] = []
        self.redo_stack: list[DraftTree] = []
        self.on_commit = on_commit
        self.draft_tree = copy.deepcopy(initial_tree) if initial_tree is not None else self._build_draft_tree()
        self.editor = DraftTreeEditor(self.draft_tree, max_wbs_depth=2)
        self.status_var = tk.StringVar(value="")
        self.type_status_var = tk.StringVar(value="")
        self.editor_state_var = tk.StringVar(value="Ready")
        self.location_var = tk.StringVar(value=f"Workbook: {self.input_path.name}    |    Sheet: {self.config.sheet_name}")
        self.detail_description_var = tk.StringVar(value="No node selected")
        self.detail_type_var = tk.StringVar(value="—")
        self.detail_level_var = tk.StringVar(value="—")
        self.detail_amount_var = tk.StringVar(value="—")
        self.detail_source_var = tk.StringVar(value="—")
        self._command_bar_mode: CommandBarMode | None = None
        self.tooltip_manager = TooltipManager(self)
        self.layout = self.workspace_services.layout
        assert self.layout is not None
        view_key = f"draft_tree::{self.input_path.resolve(strict=False)}::{self.config.sheet_name}"
        self.view_state = ViewStateManager(
            view_key,
            store=ViewStateStore(view_state_path(platform.system())),
        )
        self._register_commands()
        self._build()
        self.view_state.restore_window_geometry(self)
        self._render()

    def _register_commands(self) -> None:
        register = self.commands.register
        register(CommandID.TREE_UNDO, self._undo)
        register(CommandID.TREE_REDO, self._redo)
        register(CommandID.TREE_SAVE_DRAFT, self._save_draft)
        register(CommandID.TREE_FOCUS, self._focus_tree_shortcut)
        register(CommandID.TREE_SELECT_ALL, self._select_all_tree_shortcut)
        register(CommandID.TREE_INDENT, self._indent_shortcut)
        register(CommandID.TREE_OUTDENT, self._outdent_shortcut)
        register(CommandID.TREE_MOVE_UP, self._move_up_shortcut)
        register(CommandID.TREE_MOVE_DOWN, self._move_down_shortcut)
        register(CommandID.TREE_TOGGLE_TYPE, self._mark_wbs)
        register(CommandID.TREE_CREATE_WBS, self._create_wbs)
        register(CommandID.TREE_RENAME, self._rename)
        register(CommandID.TREE_DELETE, self._delete)
        register(CommandID.TREE_APPROVE, self._approve)
        register(CommandID.TREE_CANCEL, self._cancel)
        register(CommandID.TREE_EXPAND_SELECTED, self._expand_selected)
        register(CommandID.TREE_COLLAPSE_SELECTED, self._collapse_selected)
        register(CommandID.TREE_EXPAND_ALL, self._expand_all)
        register(CommandID.TREE_COLLAPSE_ALL, self._collapse_all)
        register(CommandID.LAYOUT_TOGGLE_SIDEBAR, self._toggle_sidebar)
        register(CommandID.LAYOUT_TOGGLE_COMMAND_BAR, self._toggle_command_bar)
        register(CommandID.LAYOUT_TOGGLE_STATUS_BAR, self._toggle_status_bar)

    def _build_draft_tree(self) -> DraftTree:
        wb = load_workbook(self.input_path, read_only=False, data_only=True)
        try:
            ws = wb[self.config.sheet_name]
            rows = list(SheetExtractor().extract(ws, self.config))
            decisions = RowClassifier().classify_rows(rows)
            return DraftTreeBuilder().build(
                self.config.sheet_name,
                rows,
                decisions,
                overrides=self.config.classification_overrides,
            )
        finally:
            wb.close()

    def _build(self) -> None:
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)
        self._build_menu_bar()
        self._build_toolbar(root)
        ttk.Label(
            root,
            text="Program detection is only a draft. Arrange WBS headings with Indent / Outdent, or create a missing WBS.",
        ).pack(fill="x", pady=(0, 8))

        self.workspace = ttk.Frame(root)
        self.workspace.pack(fill="both", expand=True)
        self.workspace.rowconfigure(0, weight=1)
        self.workspace.columnconfigure(0, weight=1)

        table = ttk.Frame(self.workspace)
        table.grid(row=0, column=0, sticky="nsew")
        self.tree = ttk.Treeview(table, columns=("type", "wbs_level", "amount", "source"), show="tree headings", selectmode="extended")
        self.tree.heading("#0", text="Description")
        self.tree.heading("type", text="Type")
        self.tree.heading("wbs_level", text="WBS Level")
        self.tree.heading("amount", text="Amount")
        self.tree.heading("source", text="Source Row")
        self.tree.column("#0", width=650, minwidth=300, stretch=True)
        self.tree.column("type", width=110, stretch=False)
        self.tree.column("wbs_level", width=100, stretch=False, anchor="center")
        self.tree.column("amount", width=130, stretch=False, anchor="e")
        self.tree.column("source", width=90, stretch=False, anchor="center")
        ybar = ttk.Scrollbar(table, orient="vertical", command=self.tree.yview)
        xbar = ttk.Scrollbar(table, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        table.rowconfigure(0, weight=1)
        table.columnconfigure(0, weight=1)

        self.sidebar = ttk.LabelFrame(self.workspace, text="NODE DETAILS", padding=10, width=self.layout.state.sidebar_width)
        self.sidebar.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        self.sidebar.grid_propagate(False)
        sidebar_header = ttk.Frame(self.sidebar)
        sidebar_header.pack(fill="x", pady=(0, 10))
        ttk.Label(sidebar_header, text="Selected node", style="StatusPrimary.TLabel").pack(side="left")
        ttk.Button(
            sidebar_header,
            text="×",
            width=3,
            command=self.commands.callback(CommandID.LAYOUT_TOGGLE_SIDEBAR),
        ).pack(side="right")
        ttk.Label(self.sidebar, textvariable=self.detail_description_var, wraplength=240, justify="left").pack(fill="x", pady=(0, 12))
        detail_grid = ttk.Frame(self.sidebar)
        detail_grid.pack(fill="x")
        details = (
            ("Type", self.detail_type_var),
            ("WBS Level", self.detail_level_var),
            ("Amount", self.detail_amount_var),
            ("Source Row", self.detail_source_var),
        )
        for row_index, (label, variable) in enumerate(details):
            ttk.Label(detail_grid, text=label, style="Hint.TLabel").grid(row=row_index, column=0, sticky="w", pady=4)
            ttk.Label(detail_grid, textvariable=variable).grid(row=row_index, column=1, sticky="e", pady=4, padx=(12, 0))
        detail_grid.columnconfigure(1, weight=1)

        self.sidebar_toggle = ttk.Button(
            self.workspace,
            text="›",
            width=3,
            command=self.commands.callback(CommandID.LAYOUT_TOGGLE_SIDEBAR),
        )
        self.sidebar_toggle.grid(row=0, column=2, sticky="ne", padx=(4, 0))

        self.tree.bind("<Double-1>", self.commands.event_callback(CommandID.TREE_RENAME))
        self.tree.bind("<Return>", self.commands.event_callback(CommandID.TREE_RENAME))
        self.tree.bind("<F2>", self.commands.event_callback(CommandID.TREE_RENAME))
        self.tree.bind("<Button-1>", self._focus_tree, add="+")
        self.tree.bind("<<TreeviewSelect>>", self._update_node_details, add="+")
        self._build_tree_context_menu()
        context_button = "<Button-2>" if is_macos(platform.system()) else "<Button-3>"
        self.tree.bind(context_button, self._show_tree_context_menu, add="+")
        if is_macos(platform.system()):
            self.tree.bind("<Control-Button-1>", self._show_tree_context_menu, add="+")
        # Native keyboard commands. Bind to this dialog rather than bind_all
        # so shortcuts disappear cleanly when the window closes.
        self.bind("<Delete>", self._on_delete_key, add="+")
        self.bind("<KP_Delete>", self._on_delete_key, add="+")
        shortcuts = shortcuts_for(platform.system())
        for sequence in shortcuts.undo:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_UNDO), add="+")
        for sequence in shortcuts.redo:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_REDO), add="+")
        for sequence in shortcuts.save:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_SAVE_DRAFT), add="+")
        for sequence in shortcuts.find:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_FOCUS), add="+")
        for sequence in shortcuts.select_all:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_SELECT_ALL, pass_event=True), add="+")
        for sequence in shortcuts.move_up:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_MOVE_UP, pass_event=True), add="+")
        for sequence in shortcuts.move_down:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_MOVE_DOWN, pass_event=True), add="+")
        for sequence in shortcuts.indent:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_INDENT, pass_event=True), add="+")
        for sequence in shortcuts.outdent:
            self.bind(sequence, self.commands.event_callback(CommandID.TREE_OUTDENT, pass_event=True), add="+")
        # Left/Right remain unbound here intentionally: ttk.Treeview already
        # provides native collapse/expand behavior and must not be overridden.
        self.bind("<Escape>", self.commands.event_callback(CommandID.TREE_CANCEL), add="+")

        self._build_command_bar(root)
        self.protocol("WM_DELETE_WINDOW", self.commands.callback(CommandID.TREE_CANCEL))

    def _add_menu_commands(self, menu: tk.Menu, specs) -> None:
        system_name = platform.system()
        for spec in specs:
            menu.add_command(
                label=spec.label,
                accelerator=normalized_shortcut(spec.shortcut, system_name=system_name),
                command=self.commands.callback(spec.command_id),
            )

    def _build_menu_bar(self) -> None:
        menu_bar = tk.Menu(self)
        file_menu = tk.Menu(menu_bar, tearoff=False)
        self._add_menu_commands(file_menu, TREE_FILE_MENU)
        menu_bar.add_cascade(label="File", menu=file_menu)

        edit_menu = tk.Menu(menu_bar, tearoff=False)
        self._add_menu_commands(edit_menu, TREE_EDIT_MENU)
        menu_bar.add_cascade(label="Edit", menu=edit_menu)

        tree_menu = tk.Menu(menu_bar, tearoff=False)
        self._add_menu_commands(tree_menu, TREE_MENU)
        menu_bar.add_cascade(label="Tree", menu=tree_menu)

        view_menu = tk.Menu(menu_bar, tearoff=False)
        self._add_menu_commands(view_menu, TREE_VIEW_MENU)
        menu_bar.add_cascade(label="View", menu=view_menu)
        self.configure(menu=menu_bar)
        self.menu_bar = menu_bar

    def _build_toolbar(self, root: ttk.Frame) -> None:
        toolbar = ttk.Frame(root, padding=(0, 0, 0, 8))
        toolbar.pack(fill="x")
        system_name = platform.system()
        groups_after = {0, 2, 6}
        for index, spec in enumerate(TREE_TOOLBAR):
            button = ttk.Button(
                toolbar,
                text=spec.label,
                command=self.commands.callback(spec.command_id),
            )
            button.pack(side="left", padx=(0, 4))
            self.tooltip_manager.register(button, tooltip_text(spec, system_name=system_name))
            if index in groups_after:
                ttk.Separator(toolbar, orient="vertical").pack(side="left", fill="y", padx=(4, 8))
        self.toolbar = toolbar

    def _build_tree_context_menu(self) -> None:
        self.tree_context_menu = tk.Menu(self, tearoff=False)
        for index, spec in enumerate(TREE_CONTEXT_MENU):
            if index in {1, 2, 6, 8, 9}:
                self.tree_context_menu.add_separator()
            self.tree_context_menu.add_command(
                label=spec.label,
                accelerator=normalized_shortcut(spec.shortcut, system_name=platform.system()),
                command=self.commands.callback(spec.command_id),
            )

    def _show_tree_context_menu(self, event) -> str:
        item = self.tree.identify_row(event.y)
        if item:
            if item not in self.tree.selection():
                self.tree.selection_set(item)
            self.tree.focus(item)
        try:
            self.tree_context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.tree_context_menu.grab_release()
        return "break"

    def _build_command_bar(self, root: ttk.Frame) -> None:
        ttk.Separator(root, orient="horizontal").pack(fill="x", pady=(10, 6))
        self.command_bar = ttk.Frame(root)
        self.command_bar.pack(fill="x")
        self.command_bar.columnconfigure(0, weight=1)
        self.command_bar.columnconfigure(1, weight=0)
        self.command_bar.columnconfigure(2, weight=0)

        self.edit_group = ttk.LabelFrame(self.command_bar, text="EDITING", padding=6)
        self.status_group = ttk.LabelFrame(self.command_bar, text="STATUS", padding=8)
        self.action_group = ttk.LabelFrame(self.command_bar, text="ACTIONS", padding=6)

        edit_commands = {
            "indent": CommandID.TREE_INDENT,
            "outdent": CommandID.TREE_OUTDENT,
            "move_up": CommandID.TREE_MOVE_UP,
            "move_down": CommandID.TREE_MOVE_DOWN,
            "toggle_type": CommandID.TREE_TOGGLE_TYPE,
            "create_wbs": CommandID.TREE_CREATE_WBS,
            "edit": CommandID.TREE_RENAME,
            "delete": CommandID.TREE_DELETE,
        }
        self.edit_buttons: list[ttk.Frame] = []
        for spec in EDIT_COMMANDS:
            cell = ttk.Frame(self.edit_group, padding=(2, 0))
            button = ttk.Button(cell, text=spec.text, command=self.commands.callback(edit_commands[spec.key]))
            button.pack(fill="x", expand=True)
            self.tooltip_manager.register(button, f"{spec.text}" + (f"\nShortcut: {spec.shortcut}" if spec.shortcut else ""))
            if spec.shortcut:
                ttk.Label(cell, text=spec.shortcut, anchor="center", style="Hint.TLabel").pack(fill="x", pady=(2, 0))
            self.edit_buttons.append(cell)

        ttk.Label(self.status_group, textvariable=self.status_var, anchor="center", style="StatusPrimary.TLabel").pack(fill="x")
        ttk.Label(self.status_group, textvariable=self.type_status_var, anchor="center", style="Hint.TLabel").pack(fill="x", pady=(4, 0))

        action_commands = {
            "save": CommandID.TREE_SAVE_DRAFT,
            "cancel": CommandID.TREE_CANCEL,
            "approve": CommandID.TREE_APPROVE,
        }
        self.action_buttons: dict[str, ttk.Frame] = {}
        for spec in ACTION_COMMANDS:
            cell = ttk.Frame(self.action_group, padding=(2, 0))
            style = "Accent.TButton" if spec.key == "approve" else "TButton"
            button = ttk.Button(cell, text=spec.text, command=self.commands.callback(action_commands[spec.key]), style=style)
            button.pack(fill="x", expand=True)
            shortcut = spec.shortcut
            self.tooltip_manager.register(button, f"{spec.text}" + (f"\nShortcut: {shortcut}" if shortcut else ""))
            if spec.key == "save" and platform.system() == "Darwin":
                shortcut = "⌘S"
            if shortcut:
                ttk.Label(cell, text=shortcut, anchor="center", style="Hint.TLabel").pack(fill="x", pady=(2, 0))
            self.action_buttons[spec.key] = cell

        self.footer = ttk.Frame(root)
        self.footer.pack(fill="x", pady=(6, 0))
        ttk.Label(self.footer, textvariable=self.editor_state_var, style="Hint.TLabel").pack(side="left")
        ttk.Label(self.footer, textvariable=self.location_var, style="Hint.TLabel").pack(side="right")

        self.layout.register_region("sidebar", GridRegion(self.sidebar))
        self.layout.register_region("command_bar", PackRegion(self.command_bar, fill="x", before=self.footer))
        self.layout.register_region("status_bar", PackRegion(self.footer, fill="x", pady=(6, 0)))
        self._sync_sidebar_controls()

        self.bind("<Configure>", self._on_dialog_configure, add="+")
        self.after_idle(lambda: self._apply_command_bar_layout(command_bar_mode(self.winfo_width())))

    def _toggle_sidebar(self) -> None:
        self.layout.toggle("sidebar")
        self._sync_sidebar_controls()

    def _toggle_command_bar(self) -> None:
        self.layout.toggle("command_bar")

    def _toggle_status_bar(self) -> None:
        self.layout.toggle("status_bar")

    def _sync_sidebar_controls(self) -> None:
        visible = self.layout.state.sidebar_visible
        self.sidebar_toggle.configure(text="‹" if visible else "›")
        self.sidebar_toggle.grid_remove() if visible else self.sidebar_toggle.grid()

    def _update_node_details(self, _event=None) -> None:
        selection = self.tree.selection()
        node = self.draft_tree.find(selection[0]) if selection else None
        if node is None:
            self.detail_description_var.set("No node selected")
            self.detail_type_var.set("—")
            self.detail_level_var.set("—")
            self.detail_amount_var.set("—")
            self.detail_source_var.set("—")
            return
        self.detail_description_var.set(node.description or "(blank description)")
        self.detail_type_var.set(node.node_type.value)
        level = self.editor.wbs_level(node.node_id)
        self.detail_level_var.set("—" if level is None else f"WBS-{level}")
        self.detail_amount_var.set("—" if node.amount is None else f"{node.amount:,.2f}")
        self.detail_source_var.set("—" if node.source_row is None else str(node.source_row))

    def _on_dialog_configure(self, event=None) -> None:
        if event is not None and getattr(event, "widget", None) is not self:
            return
        self._apply_command_bar_layout(command_bar_mode(self.winfo_width()))

    def _apply_command_bar_layout(self, mode: CommandBarMode) -> None:
        if mode == self._command_bar_mode:
            return
        self._command_bar_mode = mode
        for group in (self.edit_group, self.status_group, self.action_group):
            group.grid_forget()

        if mode == CommandBarMode.WIDE:
            self.edit_group.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
            self.status_group.grid(row=0, column=1, sticky="nsew", padx=6)
            self.action_group.grid(row=0, column=2, sticky="nsew", padx=(6, 0))
        else:
            self.edit_group.grid(row=0, column=0, columnspan=2, sticky="ew")
            self.status_group.grid(row=1, column=0, sticky="nsew", pady=(6, 0), padx=(0, 3))
            self.action_group.grid(row=1, column=1, sticky="nsew", pady=(6, 0), padx=(3, 0))

        positions = editing_grid(mode, len(self.edit_buttons))
        edit_columns = len(self.edit_buttons) if mode == CommandBarMode.WIDE else 4
        for column in range(len(self.edit_buttons)):
            self.edit_group.columnconfigure(column, weight=1 if column < edit_columns else 0)
        for cell, (row, column) in zip(self.edit_buttons, positions, strict=True):
            cell.grid(row=row, column=column, sticky="nsew")

        for column, key in enumerate(("save", "cancel", "approve")):
            self.action_group.columnconfigure(column, weight=1)
            self.action_buttons[key].grid(row=0, column=column, sticky="nsew")

    def _update_editor_state(self) -> None:
        self.editor_state_var.set("Unsaved changes" if self.dirty else "Ready")

    def _render(self, select_ids: list[str] | None = None) -> None:
        columns = ("#0", "type", "wbs_level", "amount", "source")
        if self.tree.get_children():
            self.view_state.capture(self.tree, columns, window_geometry=self.geometry())
        self.tree.delete(*self.tree.get_children())

        def insert(parent_iid: str, node) -> None:
            label = node.node_type.value
            iid = node.node_id
            self.tree.insert(
                parent_iid,
                "end",
                iid=iid,
                text=node.description,
                values=(
                    label,
                    "" if (level := self.editor.wbs_level(node.node_id)) is None else f"WBS-{level}",
                    "" if node.amount is None else f"{node.amount:,.2f}",
                    "" if node.source_row is None else node.source_row,
                ),
                open=node.is_wbs,
            )
            for child in node.children:
                insert(iid, child)

        root_iid = self.draft_tree.root.node_id
        self.tree.insert("", "end", iid=root_iid, text=self.draft_tree.sheet_name, values=("WBS-2 / Sheet", "WBS-2", "", ""), open=True)
        for child in self.draft_tree.root.children:
            insert(root_iid, child)
        summary = summarize_tree(self.draft_tree)
        self.status_var.set(summary.primary_line)
        self.type_status_var.set(summary.secondary_line)
        self._update_editor_state()
        selected_override = tuple(select_ids) if select_ids is not None else None
        self.view_state.restore(self.tree, columns, selected_override=selected_override)
        visible = [node_id for node_id in (select_ids or []) if self.tree.exists(node_id)]
        if visible:
            self.tree.see(visible[0])
        self._update_node_details()

    def _focus_tree(self, _event=None) -> None:
        self.tree.focus_set()

    def _set_open_recursive(self, item_id: str, *, opened: bool) -> None:
        if self.tree.exists(item_id):
            self.tree.item(item_id, open=opened)
            for child in self.tree.get_children(item_id):
                self._set_open_recursive(child, opened=opened)

    def _expand_selected(self) -> None:
        for item_id in self.tree.selection():
            self.tree.item(item_id, open=True)

    def _collapse_selected(self) -> None:
        for item_id in self.tree.selection():
            self.tree.item(item_id, open=False)

    def _expand_all(self) -> None:
        for item_id in self.tree.get_children(""):
            self._set_open_recursive(item_id, opened=True)

    def _collapse_all(self) -> None:
        for item_id in self.tree.get_children(""):
            self._set_open_recursive(item_id, opened=False)

    def _on_delete_key(self, event=None) -> str | None:
        # Never turn a normal text-edit Delete into a destructive tree action.
        widget = getattr(event, "widget", None)
        widget_class = widget.winfo_class() if widget is not None else None
        if is_text_input_widget(widget_class):
            return None
        if self.tree.selection():
            commands = getattr(self, "commands", None)
            if commands is None:  # Keeps the handler independently testable.
                self._delete()
            else:
                commands.execute(CommandID.TREE_DELETE)
        return "break"

    def _save_draft_shortcut(self, _event=None) -> str:
        self._save_draft()
        return "break"

    def _focus_tree_shortcut(self, _event=None) -> str:
        self.tree.focus_set()
        return "break"

    @staticmethod
    def _tree_edit_shortcut_is_safe(event=None) -> bool:
        widget = getattr(event, "widget", None)
        widget_class = widget.winfo_class() if widget is not None else None
        return not is_text_input_widget(widget_class)

    def _select_all_tree_shortcut(self, event=None) -> str | None:
        if not self._tree_edit_shortcut_is_safe(event):
            return None
        node_ids = [
            node.node_id
            for _depth, node in self.draft_tree.walk()
            if node.node_id != self.draft_tree.root.node_id
        ]
        self.tree.selection_set(node_ids)
        if node_ids:
            self.tree.focus(node_ids[0])
            self.tree.see(node_ids[0])
        self.tree.focus_set()
        return "break"

    def _move_up_shortcut(self, event=None) -> str | None:
        if not self._tree_edit_shortcut_is_safe(event):
            return None
        self._move_up()
        return "break"

    def _move_down_shortcut(self, event=None) -> str | None:
        if not self._tree_edit_shortcut_is_safe(event):
            return None
        self._move_down()
        return "break"

    def _indent_shortcut(self, event=None) -> str | None:
        if not self._tree_edit_shortcut_is_safe(event):
            return None
        self._indent()
        return "break"

    def _outdent_shortcut(self, event=None) -> str | None:
        if not self._tree_edit_shortcut_is_safe(event):
            return None
        self._outdent()
        return "break"

    def _escape(self, _event=None) -> str:
        self._cancel()
        return "break"

    def _selected_ids(self, *, require_one: bool = True) -> list[str]:
        selected = [
            node_id
            for node_id in self.tree.selection()
            if node_id != self.draft_tree.root.node_id
        ]
        if require_one and not selected:
            messagebox.showwarning("Select row", "Select one or more tree rows first.", parent=self)
        return selected

    def _selected_id(self) -> str | None:
        selected = self._selected_ids()
        if not selected:
            return None
        if len(selected) != 1:
            messagebox.showwarning("Select one row", "This action requires exactly one selected row.", parent=self)
            return None
        return selected[0]

    def _remember(self) -> None:
        self.undo_stack.append(copy.deepcopy(self.draft_tree))
        if len(self.undo_stack) > 100:
            self.undo_stack.pop(0)
        self.redo_stack.clear()

    def _restore_tree(self, tree: DraftTree) -> None:
        self.draft_tree = copy.deepcopy(tree)
        self.editor = DraftTreeEditor(self.draft_tree, max_wbs_depth=2)
        self.dirty = True
        self._render()

    def _undo(self, _event=None) -> str:
        if not self.undo_stack:
            return "break"
        self.redo_stack.append(copy.deepcopy(self.draft_tree))
        self._restore_tree(self.undo_stack.pop())
        return "break"

    def _redo(self, _event=None) -> str:
        if not self.redo_stack:
            return "break"
        self.undo_stack.append(copy.deepcopy(self.draft_tree))
        self._restore_tree(self.redo_stack.pop())
        return "break"

    def _run_edit(self, operation, node_id: str) -> None:
        self._remember()
        try:
            operation(node_id)
        except TreeEditError as exc:
            self.undo_stack.pop()
            messagebox.showwarning("Cannot edit tree", str(exc), parent=self)
            return
        self.dirty = True
        self._render(select_ids=[node_id])

    def _indent(self) -> None:
        node_id = self._selected_id()
        if node_id:
            self._run_edit(self.editor.indent, node_id)

    def _outdent(self) -> None:
        node_id = self._selected_id()
        if node_id:
            self._run_edit(self.editor.outdent, node_id)

    def _move_up(self) -> None:
        node_ids = self._selected_ids()
        if not node_ids:
            return
        self._remember()
        try:
            self.editor.move_nodes_up(node_ids)
        except TreeEditError as exc:
            self.undo_stack.pop()
            messagebox.showwarning("Cannot move rows", str(exc), parent=self)
            return
        self.dirty = True
        self._render(select_ids=node_ids)

    def _move_down(self) -> None:
        node_ids = self._selected_ids()
        if not node_ids:
            return
        self._remember()
        try:
            self.editor.move_nodes_down(node_ids)
        except TreeEditError as exc:
            self.undo_stack.pop()
            messagebox.showwarning("Cannot move rows", str(exc), parent=self)
            return
        self.dirty = True
        self._render(select_ids=node_ids)

    def _mark_wbs(self) -> None:
        node_ids = self._selected_ids()
        if not node_ids:
            return
        self._remember()
        try:
            self.editor.toggle_wbs_item(node_ids)
        except TreeEditError as exc:
            self.undo_stack.pop()
            messagebox.showwarning("Cannot toggle WBS / ITEM", str(exc), parent=self)
            return
        self.dirty = True
        self._render(select_ids=node_ids)

    def _create_wbs(self) -> None:
        selected = self.tree.selection()
        selected_id = selected[0] if len(selected) == 1 and selected[0] != self.draft_tree.root.node_id else None
        name = simpledialog.askstring("Create WBS", "WBS description", parent=self)
        if name is None:
            return
        self._remember()
        try:
            node = self.editor.create_wbs(name, selected_id)
        except TreeEditError as exc:
            self.undo_stack.pop()
            messagebox.showwarning("Cannot create WBS", str(exc), parent=self)
            return
        self.dirty = True
        self._render(select_ids=[node.node_id])

    def _rename(self) -> None:
        node_id = self._selected_id()
        if not node_id:
            return
        node = self.draft_tree.find(node_id)
        if node is None:
            return
        name = simpledialog.askstring(
            "Edit Description",
            "Description",
            initialvalue=node.description,
            parent=self,
        )
        if name is None:
            return
        self._remember()
        try:
            self.editor.rename_node(node_id, name)
        except TreeEditError as exc:
            self.undo_stack.pop()
            messagebox.showwarning("Cannot edit description", str(exc), parent=self)
            return
        self.dirty = True
        self._render(select_ids=[node_id])

    def _delete(self) -> None:
        node_ids = self._selected_ids()
        if not node_ids:
            return
        if not messagebox.askyesno(
            "Delete rows",
            f"Remove {len(node_ids)} selected row(s) from the output tree?",
            parent=self,
        ):
            return
        self._remember()
        try:
            self.editor.delete_nodes(node_ids)
        except TreeEditError as exc:
            self.undo_stack.pop()
            messagebox.showwarning("Cannot delete rows", str(exc), parent=self)
            return
        self.dirty = True
        self._render()

    def _persist_view_state(self) -> None:
        if hasattr(self, "tree") and self.tree.winfo_exists():
            columns = ("#0", "type", "wbs_level", "amount", "source")
            try:
                self.view_state.capture(self.tree, columns, window_geometry=self.geometry())
            except OSError as exc:
                messagebox.showwarning(
                    "View preferences not saved",
                    f"Could not save window preferences: {exc}\n"
                    "This does not change the saved tree or your Discard decision.",
                    parent=self,
                )

    def _finish_tree(self, mode: str) -> None:
        if self.on_commit is not None and not self.on_commit(self.draft_tree, mode):
            return  # Keep candidate and window available after failed persistence.
        self._persist_view_state()
        self.result = self.draft_tree
        self.result_mode = mode
        self.dirty = False
        self.destroy()

    def _save_draft(self) -> None:
        self._finish_tree("draft")

    def _approve(self) -> None:
        self._finish_tree("ready")

    def _cancel(self) -> None:
        if not self.dirty:
            self._persist_view_state()
            self.destroy()
            return
        answer = messagebox.askyesnocancel(
            "Unsaved tree changes",
            "Save the current tree as a draft before closing?",
            parent=self,
        )
        if answer is None:
            return
        if answer:
            self._save_draft()
        else:
            self._persist_view_state()
            self.destroy()


class BOQStudioApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("BOQ Studio — Draft Tree Export")
        self.appearance = install_appearance(self)
        self.geometry("980x720")
        self.minsize(840, 620)
        self.input_path: Path | None = None
        self.sheet_vars: dict[str, tk.BooleanVar] = {}
        self.configs: dict[str, SheetConfig] = {}
        self.events: queue.Queue = queue.Queue()
        self.draft_trees: dict[str, DraftTree] = {}
        self.tree_ready_sheets: set[str] = set()
        self.exported_sheets: set[str] = set()
        self.is_running = False
        self._session_dirty = False
        self._session_token: str | None = None
        self._source_info: dict = {}
        self.system_name = platform.system()
        self._recent_path = recent_files_path(self.system_name)
        self.recent_files = load_recent_files(self._recent_path)
        self.commands = CommandRegistry()
        self._register_commands()
        self._build()
        self._build_native_menus()
        self._bind_shortcuts()
        self.protocol("WM_DELETE_WINDOW", self.commands.callback(CommandID.FILE_CLOSE))
        self.after(100, self._poll_events)

    def _register_commands(self) -> None:
        register = self.commands.register
        register(CommandID.FILE_OPEN, self.choose_file)
        register(CommandID.FILE_SAVE_PROJECT, self._save_project_session)
        register(CommandID.FILE_CLOSE, self._request_close)
        register(CommandID.FILE_CLEAR_RECENT, self._clear_recent_files)
        register(CommandID.APP_ABOUT, self._show_about)
        register(CommandID.SHEET_FIND, self._find_shortcut)
        register(CommandID.SHEET_SELECT_ALL, self._select_all_shortcut)
        register(CommandID.SHEET_CLEAR_ALL, lambda: self._set_all(False))
        register(CommandID.SHEET_CONFIGURE, self.configure_selected)
        register(CommandID.SHEET_OPEN_DRAFT_TREE, self.open_draft_tree)
        register(CommandID.SHEET_USE, lambda: self._set_selected_sheet(True))
        register(CommandID.SHEET_IGNORE, lambda: self._set_selected_sheet(False))
        register(CommandID.WORKBOOK_NORMALIZE, self.run_normalize)

    def _build(self) -> None:
        root = ttk.Frame(self, padding=18)
        root.pack(fill="both", expand=True)
        ttk.Label(root, text="BOQ Studio", style="Title.TLabel").pack(anchor="w")
        ttk.Label(root, text="Prepare WBS-3 / WBS-4 and export an OKD-compatible workbook.", style="Subtitle.TLabel").pack(anchor="w", pady=(0, 14))

        file_frame = ttk.LabelFrame(root, text="1. Select BOQ File", padding=10)
        file_frame.pack(fill="x")
        self.file_label = ttk.Label(file_frame, text="No file selected", wraplength=700)
        self.file_label.pack(side="left", fill="x", expand=True)
        ttk.Button(file_frame, text="Choose Excel file…", command=self.commands.callback(CommandID.FILE_OPEN)).pack(side="right")

        sheets_frame = ttk.LabelFrame(root, text="2. Configure Mapping", padding=10)
        sheets_frame.pack(fill="both", expand=True, pady=12)
        toolbar = ttk.Frame(sheets_frame)
        toolbar.pack(fill="x", pady=(0, 6))
        self.select_all_button = ttk.Button(toolbar, text="Select all", command=self.commands.callback(CommandID.SHEET_SELECT_ALL), state="disabled")
        self.select_all_button.pack(side="left")
        self.clear_button = ttk.Button(toolbar, text="Clear", command=self.commands.callback(CommandID.SHEET_CLEAR_ALL), state="disabled")
        self.clear_button.pack(side="left", padx=6)
        self.configure_button = ttk.Button(toolbar, text="Configure Mapping…", command=self.commands.callback(CommandID.SHEET_CONFIGURE), state="disabled")
        self.configure_button.pack(side="left", padx=(8, 0))
        self.tree_button = ttk.Button(toolbar, text="3. Draft WBS Tree…", command=self.commands.callback(CommandID.SHEET_OPEN_DRAFT_TREE), state="disabled")
        self.tree_button.pack(side="left", padx=(8, 0))
        ttk.Label(toolbar, text="Select a sheet, configure its columns, then approve its Draft WBS Tree.").pack(side="right")

        self.sheet_tree = ttk.Treeview(sheets_frame, columns=("use", "sheet", "status"), show="headings", selectmode="browse")
        self.sheet_tree.heading("use", text="Use")
        self.sheet_tree.heading("sheet", text="Sheet")
        self.sheet_tree.heading("status", text="Status")
        self.sheet_tree.column("use", width=70, anchor="center")
        self.sheet_tree.column("sheet", width=340)
        self.sheet_tree.column("status", width=300)
        ybar = ttk.Scrollbar(sheets_frame, orient="vertical", command=self.sheet_tree.yview)
        self.sheet_tree.configure(yscrollcommand=ybar.set)
        self.sheet_tree.pack(side="left", fill="both", expand=True)
        ybar.pack(side="right", fill="y")
        self.sheet_tree.bind("<Double-1>", self.configure_selected)
        self.sheet_tree.bind("<Button-1>", self.toggle_sheet)
        self.sheet_tree.bind("<<TreeviewSelect>>", lambda _event: self._update_workflow_ui())

        action = ttk.Frame(root)
        action.pack(fill="x")
        self.run_button = ttk.Button(action, text="4. Normalize to Desktop", command=self.commands.callback(CommandID.WORKBOOK_NORMALIZE), style="Accent.TButton", state="disabled")
        self.run_button.pack(side="right")

        progress_frame = ttk.LabelFrame(root, text="Workflow status", padding=10)
        progress_frame.pack(fill="x", pady=(12, 0))
        self.status_label = ttk.Label(progress_frame, text="Step 1 — Choose a BOQ workbook")
        self.status_label.pack(anchor="w")
        self.progress = ttk.Progressbar(progress_frame, mode="determinate", maximum=100)
        self.progress.pack(fill="x", pady=(6, 0))


    def _build_native_menus(self) -> None:
        menubar = tk.Menu(self)

        file_menu = tk.Menu(menubar, tearoff=False)
        file_menu.add_command(label="Open…", command=self.commands.callback(CommandID.FILE_OPEN), accelerator="⌘O" if is_macos(self.system_name) else "Ctrl+O")
        self.recent_menu = tk.Menu(file_menu, tearoff=False)
        file_menu.add_cascade(label="Open Recent", menu=self.recent_menu)
        file_menu.add_separator()
        file_menu.add_command(label="Save Project", command=self.commands.callback(CommandID.FILE_SAVE_PROJECT), accelerator="⌘S" if is_macos(self.system_name) else "Ctrl+S")
        file_menu.add_separator()
        file_menu.add_command(label="Close Window" if is_macos(self.system_name) else "Exit", command=self.commands.callback(CommandID.FILE_CLOSE), accelerator="⌘W" if is_macos(self.system_name) else "Alt+F4")
        menubar.add_cascade(label="File", menu=file_menu)

        edit_menu = tk.Menu(menubar, tearoff=False)
        edit_menu.add_command(label="Find Sheet", command=self.commands.callback(CommandID.SHEET_FIND), accelerator="⌘F" if is_macos(self.system_name) else "Ctrl+F")
        edit_menu.add_command(label="Select All Sheets", command=self.commands.callback(CommandID.SHEET_SELECT_ALL), accelerator="⌘A" if is_macos(self.system_name) else "Ctrl+A")
        menubar.add_cascade(label="Edit", menu=edit_menu)

        if is_macos(self.system_name):
            app_menu = tk.Menu(menubar, name="apple", tearoff=False)
            app_menu.add_command(label="About BOQ Studio", command=self.commands.callback(CommandID.APP_ABOUT))
            app_menu.add_separator()
            app_menu.add_command(label="Quit BOQ Studio", command=self.commands.callback(CommandID.FILE_CLOSE), accelerator="⌘Q")
            menubar.add_cascade(menu=app_menu)
            self.createcommand("tk::mac::Quit", self.commands.callback(CommandID.FILE_CLOSE))

        help_menu = tk.Menu(menubar, tearoff=False)
        help_menu.add_command(label="About BOQ Studio", command=self.commands.callback(CommandID.APP_ABOUT))
        menubar.add_cascade(label="Help", menu=help_menu)
        self.configure(menu=menubar)
        self._refresh_recent_menu()

        self.sheet_context_menu = tk.Menu(self, tearoff=False)
        self.sheet_context_menu.add_command(label="Configure Mapping…", command=self.commands.callback(CommandID.SHEET_CONFIGURE))
        self.sheet_context_menu.add_command(label="Open Draft WBS Tree…", command=self.commands.callback(CommandID.SHEET_OPEN_DRAFT_TREE))
        self.sheet_context_menu.add_separator()
        self.sheet_context_menu.add_command(label="Use Sheet", command=self.commands.callback(CommandID.SHEET_USE))
        self.sheet_context_menu.add_command(label="Ignore Sheet", command=self.commands.callback(CommandID.SHEET_IGNORE))
        context_button = "<Button-2>" if is_macos(self.system_name) else "<Button-3>"
        self.sheet_tree.bind(context_button, self._show_sheet_context_menu, add="+")
        if is_macos(self.system_name):
            self.sheet_tree.bind("<Control-Button-1>", self._show_sheet_context_menu, add="+")

    def _refresh_recent_menu(self) -> None:
        self.recent_menu.delete(0, "end")
        self.recent_files = [path for path in self.recent_files if path.exists()]
        if not self.recent_files:
            self.recent_menu.add_command(label="No Recent Files", state="disabled")
            return
        for path in self.recent_files:
            self.recent_menu.add_command(label=path.name, command=lambda selected=path: self._open_path(selected))
        self.recent_menu.add_separator()
        self.recent_menu.add_command(label="Clear Menu", command=self.commands.callback(CommandID.FILE_CLEAR_RECENT))

    def _remember_recent_file(self, path: Path) -> None:
        self.recent_files = add_recent_file(path, self.recent_files)
        try:
            save_recent_files(self._recent_path, self.recent_files)
        except OSError:
            pass
        self._refresh_recent_menu()

    def _clear_recent_files(self) -> None:
        self.recent_files = []
        try:
            save_recent_files(self._recent_path, [])
        except OSError:
            pass
        self._refresh_recent_menu()

    def _open_path(self, path: Path) -> None:
        if not path.exists():
            messagebox.showwarning("File unavailable", f"The recent workbook no longer exists:\n{path}", parent=self)
            self.recent_files = [item for item in self.recent_files if item != path]
            self._refresh_recent_menu()
            return
        self._load_workbook_path(path)

    def _show_sheet_context_menu(self, event) -> str:
        row = self.sheet_tree.identify_row(event.y)
        if row:
            self.sheet_tree.selection_set(row)
            self.sheet_tree.focus(row)
            self._update_workflow_ui()
            self.sheet_context_menu.tk_popup(event.x_root, event.y_root)
        return "break"

    def _set_selected_sheet(self, value: bool) -> None:
        if self._busy():
            return
        selection = self.sheet_tree.selection()
        if not selection:
            return
        name = selection[0]
        self.sheet_vars[name].set(value)
        self._refresh_row(name)
        self._update_workflow_ui()
        self._session_dirty = True
        self._save_project_session()

    def _show_about(self) -> None:
        messagebox.showinfo("About BOQ Studio", "BOQ Studio\nPrepare WBS trees and export OKD-compatible workbooks.", parent=self)

    def _request_close(self) -> None:
        if self._busy():
            return
        if self._leave_current():
            self.destroy()

    def _busy(self) -> bool:
        if self.is_running:
            messagebox.showinfo("Export in progress", "Wait for the export to finish before changing or closing this work.", parent=self)
            return True
        return False

    def _leave_current(self) -> bool:
        if not self._session_dirty or self.input_path is None:
            return True
        if self._save_project_session():
            return True
        if not messagebox.askyesno(
            "Keep recoverable work",
            "The session was not saved. Save a separate recovery JSON before leaving?\n"
            "No keeps this work open. Recovery needs the unchanged original workbook.", parent=self,
        ):
            return False
        return self._write_recovery_copy(self._session_payload())

    def _write_recovery_copy(self, payload: dict) -> bool:
        filename = filedialog.asksaveasfilename(
            parent=self, title="Save recovery JSON (choose a new filename)",
            initialfile=f"{self.input_path.stem}_recovery.json", defaultextension=".json",
            filetypes=[("Session recovery", "*.json")],
        )
        if not filename:
            return False
        try:
            target = Path(filename)
            if target.exists():
                raise ValueError("Choose a new filename; existing recovery files are not overwritten.")
            save_session(target, payload, expected_token=None)
        except Exception as exc:
            messagebox.showerror("Recovery not saved", str(exc), parent=self)
            return False
        messagebox.showinfo("Recovery saved", f"Work preserved at:\n{target}\nKeep the original workbook unchanged. See the S0 recovery guide to resume.", parent=self)
        return True

    def _bind_shortcuts(self) -> None:
        shortcuts = shortcuts_for(platform.system())
        for sequence in shortcuts.open_file:
            self.bind(sequence, self.commands.event_callback(CommandID.FILE_OPEN), add="+")
        for sequence in shortcuts.save:
            self.bind(sequence, self.commands.event_callback(CommandID.FILE_SAVE_PROJECT), add="+")
        for sequence in shortcuts.find:
            self.bind(sequence, self.commands.event_callback(CommandID.SHEET_FIND), add="+")
        for sequence in shortcuts.select_all:
            self.bind(sequence, self.commands.event_callback(CommandID.SHEET_SELECT_ALL, pass_event=True), add="+")

    def _open_shortcut(self, _event=None) -> str:
        self.choose_file()
        return "break"

    def _save_shortcut(self, _event=None) -> str:
        self._save_project_session()
        return "break"

    def _find_shortcut(self, _event=None) -> str:
        if self.sheet_tree.get_children():
            self.sheet_tree.focus_set()
        return "break"

    def _select_all_shortcut(self, event=None) -> str | None:
        widget = getattr(event, "widget", None)
        widget_class = widget.winfo_class() if widget is not None else None
        if is_text_input_widget(widget_class):
            return None
        if self.input_path is not None:
            self._set_all(True)
        return "break"

    def choose_file(self) -> None:
        filename = filedialog.askopenfilename(title="Choose BOQ Excel file", filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")])
        if not filename:
            return
        self._load_workbook_path(Path(filename))

    def _load_workbook_path(self, path: Path) -> None:
        if self._busy():
            return
        path = path.resolve()
        try:
            fingerprint = workbook_fingerprint(path)
            wb = load_workbook(path, read_only=True, data_only=False)
            try:
                names = wb.sheetnames
            finally:
                wb.close()
        except Exception as exc:
            messagebox.showerror("Cannot open workbook", str(exc), parent=self)
            return
        if not self._leave_current():
            return
        try:
            prepared = self._prepare_session(path, names, fingerprint)
            if prepared is None:
                return
            payload, token, needs_save = prepared
            configs = {name: config_from_dict(data) for name, data in payload["configs"].items()} if payload else {}
            trees = {name: tree_from_dict(data) for name, data in payload["trees"].items()} if payload else {}
            selected = dict(payload["selected"]) if payload else {}
            variables = {}
            for name in names:
                default = not any(word in name.casefold() for word in ("sum", "summary", "ปร.6", "ปร.5", "log"))
                variables[name] = tk.BooleanVar(value=selected.get(name, default))
                # A resumed session is authoritative; templates cannot fill/replace it.
                if payload is None:
                    template = default_template_path(path, name)
                    if template.exists():
                        try:
                            cfg = load_template(template, sheet_name=name)
                            configs[name] = replace(cfg, classification_overrides={}, wbs_level_overrides={})
                        except Exception:
                            pass
            if workbook_fingerprint(path) != fingerprint:
                raise ValueError("Workbook changed while opening. Save/close it in Excel and try again.")
            if session_token(session_path_for(path)) != token:
                raise ValueError("Session changed in another window while opening. Try again.")
            stat = path.stat()
        except Exception as exc:
            messagebox.showerror("Cannot open session", str(exc), parent=self)
            return
        # Commit only after all file I/O, decoding and candidate construction succeeded.
        self.input_path = path
        self._source_info = {"path": str(path), "sha256": fingerprint, "size": stat.st_size, "modified_ns": stat.st_mtime_ns}
        self._session_token = token
        self.sheet_vars = variables
        self.configs = configs
        self.draft_trees = trees
        self.tree_ready_sheets = set(payload["tree_ready"]) if payload else set()
        self.exported_sheets = set(payload["exported"]) if payload else set()
        self._session_dirty = needs_save
        self.file_label.configure(text=str(path))
        self.sheet_tree.delete(*self.sheet_tree.get_children())
        for name in names:
            self.sheet_tree.insert("", "end", iid=name, values=("✓" if variables[name].get() else "", name, self._state_for(name).value))
        self.select_all_button.configure(state="normal")
        self.clear_button.configure(state="normal")
        self._update_workflow_ui()
        try:
            self._remember_recent_file(path)
        except Exception as exc:
            messagebox.showwarning("Recent files not saved", str(exc), parent=self)

    def _prepare_session(self, workbook: Path, names: list[str], fingerprint: str):
        """Return a staged payload/token, or None on Cancel. No active state mutation."""
        path = session_path_for(workbook)
        token = session_token(path)
        backup = path.with_suffix(path.suffix + ".bak")
        if token is None and not backup.exists():
            return None, None, True
        payload = None
        problem = "Saved session is missing."
        try:
            payload = load_session(path)
            if payload["source"]["sha256"] != fingerprint:
                raise ValueError("The source workbook changed; old row decisions cannot be reused safely.")
            used = set(payload["selected"]) | set(payload["configs"]) | set(payload["trees"])
            if not used <= set(names):
                raise ValueError("Saved session refers to missing worksheets.")
        except Exception as exc:
            problem = str(exc)
            payload = None
        if payload is not None:
            answer = messagebox.askyesnocancel(
                "Resume project",
                "Resume the saved mappings and tree?\nYes: Resume. No: keep a recovery copy and start fresh. Cancel: keep current work.", parent=self,
            )
            if answer is None:
                return None
            if answer:
                return payload, token, False
        else:
            try:
                recovered = load_session(backup)
                if recovered["source"]["sha256"] != fingerprint:
                    raise ValueError("Backup is for another source revision.")
                if not (set(recovered["selected"]) | set(recovered["configs"]) | set(recovered["trees"])) <= set(names):
                    raise ValueError("Backup worksheets do not match.")
            except Exception:
                recovered = None
            if recovered is not None:
                answer = messagebox.askyesnocancel(
                    "Session recovery",
                    f"{problem}\nOpen the previous saved backup?\nYes: recover backup. No: start fresh with a recovery copy. Cancel: keep current work.", parent=self,
                )
                if answer is None:
                    return None
                if answer:
                    if token is not None:
                        archive_session(path, expected_token=token)
                    return recovered, token, True
            elif not messagebox.askyesno(
                "Saved work cannot be resumed",
                f"{problem}\nKeep a recovery copy and start fresh?\nNo cancels opening; existing work is not replaced.", parent=self,
            ):
                return None
        if token is not None:
            saved = archive_session(path, expected_token=token)
            messagebox.showinfo("Previous session preserved", f"Recovery copy:\n{saved}\nKeep the matching original workbook to restore it.", parent=self)
        return None, token, True

    def _session_payload(self) -> dict:
        return {
            "version": SESSION_VERSION, "saved_at": datetime.now(timezone.utc).isoformat(),
            "source": dict(self._source_info),
            "selected": {name: var.get() for name, var in self.sheet_vars.items()},
            "configs": {name: config_to_dict(cfg) for name, cfg in self.configs.items()},
            "trees": {name: tree_to_dict(tree) for name, tree in self.draft_trees.items()},
            "tree_ready": sorted(self.tree_ready_sheets), "exported": sorted(self.exported_sheets),
        }

    def _source_unchanged(self) -> bool:
        try:
            if self.input_path is None or workbook_fingerprint(self.input_path) != self._source_info.get("sha256"):
                raise ValueError("Source workbook changed. Restore the unchanged original to continue this session, or create new work from the revised workbook.")
        except Exception as exc:
            messagebox.showwarning("Source unavailable or changed", str(exc), parent=self)
            return False
        return True

    def _save_project_session(self) -> bool:
        if self.input_path is None:
            return True
        if self._busy():
            return False
        if not self._source_unchanged():
            return False
        try:
            payload = self._session_payload()
            self._session_token = save_session(session_path_for(self.input_path), payload, expected_token=self._session_token)
        except Exception as exc:
            self._session_dirty = True
            messagebox.showwarning("Project save failed", f"{exc}\nWork remains open and unsaved. Retry Save; closing can save a separate recovery JSON.", parent=self)
            self.status_label.configure(text="Session not saved — work remains open")
            return False
        self._session_dirty = False
        self.status_label.configure(text="Session saved")
        return True

    def _set_all(self, value: bool) -> None:
        if self._busy():
            return
        for name, var in self.sheet_vars.items():
            var.set(value)
            self._refresh_row(name)
        self._update_workflow_ui()
        self._session_dirty = True
        self._save_project_session()

    def toggle_sheet(self, event) -> None:
        if self._busy():
            return
        region = self.sheet_tree.identify("region", event.x, event.y)
        column = self.sheet_tree.identify_column(event.x)
        row = self.sheet_tree.identify_row(event.y)
        if region == "cell" and column == "#1" and row:
            self.sheet_vars[row].set(not self.sheet_vars[row].get())
            self._refresh_row(row)
            self._update_workflow_ui()
            self._session_dirty = True
            self._save_project_session()

    def configure_selected(self, _event=None) -> None:
        if self._busy() or not self._source_unchanged():
            return
        selection = self.sheet_tree.selection()
        if not selection or not self.input_path:
            return
        name = selection[0]
        previous = self.configs.get(name)
        dialog = SheetMappingDialog(self, self.input_path, name, copy.deepcopy(previous))
        self.wait_window(dialog)
        if dialog.result is None or dialog.result == previous:
            return
        if not self._source_unchanged():
            return
        if name in self.draft_trees:
            if not messagebox.askyesno(
                "Mapping changes require a new tree",
                "This changes the interpretation of source rows. Keep a recovery copy, remove this sheet's tree and approval, and apply the new mapping?\nNo keeps the reviewed tree and mapping.", parent=self,
            ):
                return
            if not self._save_project_session():
                return
            try:
                saved = archive_session(session_path_for(self.input_path), expected_token=self._session_token)
            except Exception as exc:
                messagebox.showerror("Cannot preserve reviewed tree", str(exc), parent=self)
                return
            messagebox.showinfo("Reviewed tree preserved", f"Previous mapping and tree saved at:\n{saved}", parent=self)
        self.configs[name] = dialog.result
        self.draft_trees.pop(name, None)
        self.tree_ready_sheets.discard(name)
        self.exported_sheets.discard(name)
        self.sheet_vars[name].set(True)
        self._session_dirty = True
        self._refresh_row(name)
        self._update_workflow_ui()
        if self._save_project_session():
            try:
                save_template(default_template_path(self.input_path, name), dialog.result)
            except Exception as exc:
                messagebox.showwarning("Template not saved", f"Session saved, but mapping template could not be written:\n{exc}", parent=self)

    def _commit_tree(self, name: str, tree: DraftTree, mode: str) -> bool:
        old_trees = self.draft_trees.copy()
        old_ready, old_exported = self.tree_ready_sheets.copy(), self.exported_sheets.copy()
        old_dirty = self._session_dirty
        self.draft_trees[name] = copy.deepcopy(tree)
        if mode == "ready":
            self.tree_ready_sheets.add(name)
        else:
            self.tree_ready_sheets.discard(name)
        self.exported_sheets.discard(name)
        self._session_dirty = True
        if not self._save_project_session():
            candidate_payload = self._session_payload()
            self.draft_trees = old_trees
            self.tree_ready_sheets, self.exported_sheets = old_ready, old_exported
            self._session_dirty = old_dirty
            if messagebox.askyesno(
                "Tree remains unsaved",
                "Keep a separate recovery JSON of this candidate tree?\n"
                "The editor stays open either way; No lets you retry saving.", parent=self,
            ):
                self._write_recovery_copy(candidate_payload)
            return False
        self._refresh_row(name)
        self._update_workflow_ui()
        return True

    def open_draft_tree(self) -> None:
        if self._busy() or not self._source_unchanged():
            return
        selection = self.sheet_tree.selection()
        if not selection or not self.input_path:
            messagebox.showwarning("Select sheet", "Select one configured sheet first.", parent=self)
            return
        name = selection[0]
        cfg = self.configs.get(name)
        if cfg is None:
            messagebox.showwarning("Mapping required", "Double-click the sheet and save its column mapping first.", parent=self)
            return
        dialog = DraftTreeDialog(
            self, self.input_path, cfg, self.draft_trees.get(name),
            on_commit=lambda tree, mode: self._commit_tree(name, tree, mode),
        )
        self.wait_window(dialog)

    def _state_for(self, name: str) -> SheetState:
        return sheet_state(
            has_mapping=name in self.configs,
            has_tree=name in self.draft_trees,
            tree_ready=name in self.tree_ready_sheets,
            exported=name in self.exported_sheets,
        )

    def _refresh_row(self, name: str) -> None:
        state = self._state_for(name)
        self.sheet_tree.item(
            name,
            values=("✓" if self.sheet_vars[name].get() else "", name, state.value),
        )

    def _selected_sheet_names(self) -> list[str]:
        return [name for name, var in self.sheet_vars.items() if var.get()]

    def _update_workflow_ui(self) -> None:
        selected_row = self.sheet_tree.selection()
        current_name = selected_row[0] if selected_row else None
        has_file = self.input_path is not None
        self.configure_button.configure(state="normal" if has_file and current_name else "disabled")
        can_open_tree = bool(current_name and current_name in self.configs)
        self.tree_button.configure(state="normal" if can_open_tree else "disabled")

        selected = self._selected_sheet_names()
        status = workflow_status(self._state_for(name) for name in selected)
        self.run_button.configure(
            state="normal" if status.can_export and not self.is_running else "disabled"
        )
        if not has_file:
            text = "Step 1 — Choose a BOQ workbook"
        elif status.selected_count == 0:
            text = "Select at least one breakdown sheet"
        elif status.new_count:
            text = f"Step 2 — Configure Mapping · {status.summary}"
        elif status.mapped_count or status.tree_draft_count:
            text = f"Step 3 — Review Draft WBS Tree · {status.summary}"
        else:
            text = f"Step 4 — {status.summary}"
        if self._session_dirty:
            text += " · Unsaved session changes"
        self.status_label.configure(text=text)

    def run_normalize(self) -> None:
        if self._busy() or not self._source_unchanged():
            return
        if not self.input_path:
            return
        selected = self._selected_sheet_names()
        if not selected:
            messagebox.showwarning("No sheets", "Select at least one breakdown sheet.", parent=self)
            return
        missing = [name for name in selected if name not in self.configs]
        if missing:
            messagebox.showwarning("Mapping required", "Configure every selected sheet first:\n\n" + "\n".join(missing), parent=self)
            return
        missing_trees = [name for name in selected if name not in self.tree_ready_sheets or name not in self.draft_trees]
        if missing_trees:
            messagebox.showwarning(
                "Draft WBS Tree required",
                "Open Draft WBS Tree and click 'Use this tree' for every selected sheet:\n\n"
                + "\n".join(missing_trees),
                parent=self,
            )
            return

        if not self._save_project_session():
            return

        desktop = Path.home() / "Desktop"
        stem = self.input_path.stem.strip()
        output = desktop / f"{stem}_OKD.xlsx"
        suffix = 2
        while output.exists():
            output = desktop / f"{stem}_OKD_{suffix}.xlsx"
            suffix += 1
        request = NormalizeRequest(
            self.input_path,
            output,
            copy.deepcopy([self.configs[name] for name in selected]),
            draft_trees=copy.deepcopy({name: self.draft_trees[name] for name in selected}),
        )
        expected_source = self._source_info["sha256"]
        self.is_running = True
        self._update_workflow_ui()
        self.progress["value"] = 0
        self.status_label.configure(text="Starting…")

        def worker() -> None:
            try:
                def on_progress(done: int, total: int, message: str) -> None:
                    self.events.put(("progress", done, total, message))
                with tempfile.TemporaryDirectory(prefix="boq-export-") as directory:
                    snapshot = Path(directory) / request.input_path.name
                    shutil.copyfile(request.input_path, snapshot)
                    if workbook_fingerprint(snapshot) != expected_source:
                        raise ValueError("Source changed before export. No output was published.")
                    # Read from snapshot but preserve original source provenance in LOG.
                    results = normalize_workbook(request, on_progress, source_snapshot=snapshot)
                self.events.put(("done", output, results, tuple(selected)))
            except Exception as exc:
                self.events.put(("error", str(exc)))

        threading.Thread(target=worker, daemon=True).start()

    def _poll_events(self) -> None:
        try:
            while True:
                event = self.events.get_nowait()
                if event[0] == "progress":
                    _, done, total, message = event
                    self.progress["value"] = (done / max(1, total)) * 100
                    self.status_label.configure(text=message)
                elif event[0] == "done":
                    _, output, results, exported_selection = event
                    self.progress["value"] = 100
                    ok = sum(1 for r in results if r.status == "OK")
                    warning = sum(1 for r in results if r.status == "WARNING")
                    error = sum(1 for r in results if r.status == "ERROR")
                    self.status_label.configure(text=f"Finished — OK {ok}, Warning {warning}, Error {error}")
                    self.exported_sheets.difference_update(exported_selection)
                    self.exported_sheets.update(r.sheet_name for r in results if r.status == "OK" and r.sheet_name in exported_selection)
                    self._session_dirty = True
                    for name in exported_selection:
                        if name in self.sheet_vars:
                            self._refresh_row(name)
                    self.is_running = False
                    self._update_workflow_ui()
                    saved = self._save_project_session()
                    if not saved:
                        self.status_label.configure(text="Excel exported — session status NOT saved")
                    messagebox.showinfo("BOQ Studio finished", f"Output created:\n{output}\n\nOK: {ok}\nWarnings: {warning}\nErrors: {error}\nSession status: {'saved' if saved else 'NOT saved; work remains open'}", parent=self)
                elif event[0] == "error":
                    self.is_running = False
                    self.status_label.configure(text="Failed")
                    self._update_workflow_ui()
                    messagebox.showerror("Normalization failed", event[1], parent=self)
        except queue.Empty:
            pass
        self.after(100, self._poll_events)


def main() -> None:
    enable_windows_dpi_awareness()
    app = BOQStudioApp()
    app.mainloop()


if __name__ == "__main__":
    main()
