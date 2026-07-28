from __future__ import annotations

import math
import re
from collections.abc import Callable
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import NormalizeRequest, SheetConfig, SheetResult

ProgressCallback = Callable[[int, int, str], None]

OUTPUT_HEADERS = [
    "WBS Code",
    "WBS Level",
    "Node Type",
    "WBS-1",
    "WBS-2",
    "WBS-3",
    "WBS-4",
    "Description",
    "Unit",
    "Qty",
    "Material",
    "Labor",
    "Amount",
    "Source Sheet",
    "Source Row",
]

TOTAL_WORDS = (
    "รวม",
    "รวมเป็นเงิน",
    "รวมทั้งสิ้น",
    "ยอดรวม",
    "subtotal",
    "grand total",
    "total",
)
NOTE_WORDS = ("หมายเหตุ", "note", "remark")


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.replace("\n", " ").split())
    return str(value).strip()


def safe_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return float(value)
    text = safe_text(value).replace(",", "")
    text = text.replace("฿", "").replace("บาท", "").strip()
    if text in {"-", "–", "—"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def numbering_level(text: str) -> int | None:
    text = text.strip()
    match = re.match(r"^(\d+(?:\.\d+){0,3})[.)\-\s]", text)
    if not match:
        match = re.fullmatch(r"(\d+(?:\.\d+){0,3})", text)
    if match:
        return min(4, match.group(1).count(".") + 1)
    thai_match = re.match(r"^(?:หมวด|ส่วนที่|งานที่)\s*(\d+)", text, flags=re.I)
    if thai_match:
        return 1
    return None


def is_total(description: str) -> bool:
    lowered = description.casefold()
    return any(word in lowered for word in TOTAL_WORDS)


def is_note(description: str) -> bool:
    lowered = description.casefold()
    return any(lowered.startswith(word) for word in NOTE_WORDS)


def cell_is_bold(ws, row: int, cols: list[int]) -> bool:
    return any(bool(ws.cell(row=row, column=col).font.bold) for col in cols)


def description_indent_level(ws, row: int, cols: list[int]) -> int:
    for index, col in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=col)
        if safe_text(cell.value):
            indent = int(cell.alignment.indent or 0)
            return min(4, max(1, index + indent))
    return 1


def classify_row(ws, row: int, cfg: SheetConfig, description: str) -> str:
    if not description:
        return "EMPTY"
    if is_total(description):
        return "TOTAL"
    if is_note(description):
        return "NOTE"

    numeric_values = [
        safe_number(ws.cell(row, col).value) if col else None
        for col in (cfg.qty_col, cfg.material_col, cfg.labor_col, cfg.amount_col)
    ]
    unit = safe_text(ws.cell(row, cfg.unit_col).value) if cfg.unit_col else ""
    if unit or any(value is not None for value in numeric_values):
        return "ITEM"

    item_no = safe_text(ws.cell(row, cfg.item_no_col).value) if cfg.item_no_col else ""
    if numbering_level(item_no) or numbering_level(description):
        return "WBS"
    if cell_is_bold(ws, row, cfg.description_cols):
        return "WBS"
    return "WBS"


def get_wbs_level(ws, row: int, cfg: SheetConfig, description: str) -> int:
    item_no = safe_text(ws.cell(row, cfg.item_no_col).value) if cfg.item_no_col else ""
    return min(
        cfg.max_wbs_level,
        numbering_level(item_no)
        or numbering_level(description)
        or description_indent_level(ws, row, cfg.description_cols),
    )


def normalize_sheet(source_ws, output_ws, cfg: SheetConfig, progress: ProgressCallback | None = None) -> SheetResult:
    result = SheetResult(sheet_name=cfg.sheet_name, status="OK")
    counters = [0] * cfg.max_wbs_level
    current_names = [""] * cfg.max_wbs_level
    max_row = source_ws.max_row
    total_rows = max(0, max_row - cfg.start_row + 1)

    output_ws.append(OUTPUT_HEADERS)
    for cell in output_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, source_row in enumerate(range(cfg.start_row, max_row + 1), start=1):
        result.source_rows += 1
        description_parts = [safe_text(source_ws.cell(source_row, col).value) for col in cfg.description_cols]
        description = " ".join(part for part in description_parts if part)
        row_type = classify_row(source_ws, source_row, cfg, description)

        if row_type in {"EMPTY", "TOTAL", "NOTE"}:
            result.skipped_rows += 1
            if progress:
                progress(index, total_rows, f"{cfg.sheet_name}: row {source_row}")
            continue

        if row_type == "WBS":
            level = get_wbs_level(source_ws, source_row, cfg, description)
            counters[level - 1] += 1
            for pos in range(level, cfg.max_wbs_level):
                counters[pos] = 0
                current_names[pos] = ""
            current_names[level - 1] = description
            code = ".".join(str(n) for n in counters[:level] if n > 0)
            output_ws.append([
                code,
                level,
                "WBS",
                *current_names[:4],
                description,
                "",
                None,
                None,
                None,
                None,
                cfg.sheet_name,
                source_row,
            ])
            result.wbs_rows += 1
        else:
            active_levels = [idx + 1 for idx, name in enumerate(current_names) if name]
            level = max(active_levels, default=1)
            code = ".".join(str(n) for n in counters[:level] if n > 0)
            output_ws.append([
                code,
                level,
                "ITEM",
                *current_names[:4],
                description,
                safe_text(source_ws.cell(source_row, cfg.unit_col).value) if cfg.unit_col else "",
                safe_number(source_ws.cell(source_row, cfg.qty_col).value) if cfg.qty_col else None,
                safe_number(source_ws.cell(source_row, cfg.material_col).value) if cfg.material_col else None,
                safe_number(source_ws.cell(source_row, cfg.labor_col).value) if cfg.labor_col else None,
                safe_number(source_ws.cell(source_row, cfg.amount_col).value) if cfg.amount_col else None,
                cfg.sheet_name,
                source_row,
            ])
            result.item_rows += 1

        result.output_rows += 1
        if progress:
            progress(index, total_rows, f"{cfg.sheet_name}: row {source_row}")

    style_output_sheet(output_ws)
    if result.output_rows == 0:
        result.status = "WARNING"
        result.message = "No WBS or item rows were produced. Check the mapping."
    return result


def style_output_sheet(ws) -> None:
    widths = {
        "A": 14, "B": 11, "C": 12,
        "D": 28, "E": 28, "F": 28, "G": 28,
        "H": 48, "I": 12,
        "J": 14, "K": 16, "L": 16, "M": 16,
        "N": 24, "O": 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows(min_row=2):
        node_type = row[2].value
        if node_type == "WBS":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(bold=True)
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in (10, 11, 12, 13):
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            cell[0].number_format = '#,##0.00;[Red]-#,##0.00;'


def normalize_workbook(request: NormalizeRequest, progress: ProgressCallback | None = None) -> list[SheetResult]:
    source_wb = load_workbook(request.input_path, data_only=True, read_only=False)
    output_wb = Workbook()
    output_wb.remove(output_wb.active)
    results: list[SheetResult] = []

    total_sheets = len(request.sheets)
    for sheet_index, cfg in enumerate(request.sheets, start=1):
        if cfg.sheet_name not in source_wb.sheetnames:
            results.append(SheetResult(cfg.sheet_name, "ERROR", message="Sheet not found"))
            continue
        output_name = cfg.sheet_name[:31]
        if output_name in output_wb.sheetnames:
            suffix = 2
            base = output_name[:27]
            while f"{base}_{suffix}" in output_wb.sheetnames:
                suffix += 1
            output_name = f"{base}_{suffix}"
        target_ws = output_wb.create_sheet(output_name)

        def sheet_progress(done: int, total: int, message: str) -> None:
            if progress:
                overall_done = (sheet_index - 1) + (done / max(1, total))
                progress(int(overall_done * 1000), total_sheets * 1000, message)

        try:
            results.append(normalize_sheet(source_wb[cfg.sheet_name], target_ws, cfg, sheet_progress))
        except Exception as exc:  # keep other sheets running
            output_wb.remove(target_ws)
            results.append(SheetResult(cfg.sheet_name, "ERROR", message=str(exc)))

    add_log_sheet(output_wb, request, results)
    request.output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(request.output_path)
    source_wb.close()
    output_wb.close()
    if progress:
        progress(1, 1, f"Saved: {request.output_path}")
    return results


def add_log_sheet(wb, request: NormalizeRequest, results: list[SheetResult]) -> None:
    ws = wb.create_sheet("LOG", 0)
    ws.append(["BOQ Studio MS-1", "Normalization Log"])
    ws.append(["Input", str(request.input_path)])
    ws.append(["Output", str(request.output_path)])
    ws.append([])
    ws.append(["Sheet", "Status", "Source Rows", "Output Rows", "WBS Rows", "Item Rows", "Skipped Rows", "Message"])
    for result in results:
        ws.append([
            result.sheet_name,
            result.status,
            result.source_rows,
            result.output_rows,
            result.wbs_rows,
            result.item_rows,
            result.skipped_rows,
            result.message,
        ])
    for cell in ws[1]:
        cell.font = Font(size=14, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for cell in ws[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "A6"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    for col in range(3, 8):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["H"].width = 50
