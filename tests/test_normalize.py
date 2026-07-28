from pathlib import Path

from openpyxl import Workbook, load_workbook

from boq_studio.models import NormalizeRequest, SheetConfig
from boq_studio.normalize import normalize_workbook


def test_normalize_filters_zero_amount_and_builds_okd_wbs(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "source_OKD.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Breakdown"
    ws.append(["No", "Description", "Unit", "Qty", "Material", "Labor", "Amount"])
    ws.append([1, "งานโครงสร้าง", None, None, None, None, None])
    ws.append(["1.1", "งานคอนกรีต", None, None, None, None, None])
    ws.append([1, "คอนกรีต 240 ksc", "m3", 10, 20000, 5000, 25000])
    ws.append([2, "คอนกรีตไม่มีราคา", "m3", 5, 0, 0, 0])
    ws.append([3, "รายการไม่มี Amount", "m3", 2, 100, 50, None])
    ws.append([None, "รวม", None, None, None, None, 25000])
    wb.save(source)

    request = NormalizeRequest(
        input_path=source,
        output_path=output,
        sheets=[SheetConfig(
            sheet_name="Breakdown",
            header_row=1,
            start_row=2,
            item_no_col=1,
            description_cols=[2],
            unit_col=3,
            qty_col=4,
            material_col=5,
            labor_col=6,
            amount_col=7,
        )],
    )
    results = normalize_workbook(request)
    assert results[0].output_rows == 1
    assert results[0].item_rows == 1
    assert results[0].wbs_rows == 2

    out = load_workbook(output, data_only=True)
    ws_out = out["Breakdown"]
    assert ws_out.max_row == 2
    assert ws_out["A2"].value == "1.1.1"
    assert ws_out["D2"].value == "Project"
    assert ws_out["E2"].value == "Breakdown"
    assert ws_out["F2"].value == "งานโครงสร้าง"
    assert ws_out["G2"].value == "งานคอนกรีต"
    assert ws_out["M2"].value == 25000


def test_project_sheet_combines_normalized_sheets_in_selected_order(tmp_path: Path):
    source = tmp_path / "multi.xlsx"
    output = tmp_path / "multi_OKD.xlsx"
    wb = Workbook()
    first = wb.active
    first.title = "Arch"
    second = wb.create_sheet("Structure")
    for ws, description, amount in (
        (first, "งานสถาปัตย์", 1000),
        (second, "งานโครงสร้าง", 2000),
    ):
        ws.append(["No", "Description", "Unit", "Qty", "Material", "Labor", "Amount"])
        ws.append([1, description, "งาน", 1, amount, 0, amount])
    wb.save(source)

    configs = [
        SheetConfig(
            sheet_name=name,
            header_row=1,
            start_row=2,
            item_no_col=1,
            description_cols=[2],
            unit_col=3,
            qty_col=4,
            material_col=5,
            labor_col=6,
            amount_col=7,
        )
        for name in ("Arch", "Structure")
    ]
    normalize_workbook(NormalizeRequest(source, output, configs))

    out = load_workbook(output, data_only=True)
    assert out.sheetnames[:2] == ["LOG", "Project"]
    project = out["Project"]
    assert project.max_row == 3
    assert project["E2"].value == "Arch"
    assert project["H2"].value == "งานสถาปัตย์"
    assert project["M2"].value == 1000
    assert project["E3"].value == "Structure"
    assert project["H3"].value == "งานโครงสร้าง"
    assert project["M3"].value == 2000
    assert out["LOG"]["B4"].value == 2
