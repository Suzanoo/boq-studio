from openpyxl import Workbook

from boq_studio.core.classification import RowClassifier, RowType
from boq_studio.core.extraction import SheetExtractor
from boq_studio.core.wbs import WBSBuilder
from boq_studio.models import SheetConfig


def config() -> SheetConfig:
    return SheetConfig("BOQ", 1, 2, 1, [2], 3, 4, 5, 6, 7)


def test_structural_total_detection_is_language_independent():
    wb = Workbook(); ws = wb.active
    ws.append(["No", "Description", "Unit", "Qty", "Material", "Labor", "Amount"])
    ws.append([None, "Any language summary", None, None, None, None, 5000])
    row = SheetExtractor().extract_row(ws, 2, config())
    assert RowClassifier().classify(row) == RowType.TOTAL


def test_item_with_total_word_is_preserved():
    wb = Workbook(); ws = wb.active
    ws.append(["No", "Description", "Unit", "Qty", "Material", "Labor", "Amount"])
    ws.append([1, "AIR SPLIT TYPE total installation accessories", "set", 1, 100, 20, 120])
    row = SheetExtractor().extract_row(ws, 2, config())
    assert RowClassifier().classify(row) == RowType.ITEM


def test_classifier_and_wbs_builder_are_independent():
    wb = Workbook(); ws = wb.active
    ws.append(["No", "Description", "Unit", "Qty", "Material", "Labor", "Amount"])
    ws.append(["1", "Structural Work", None, None, None, None, None])
    ws.append(["1.1", "Concrete", None, None, None, None, None])
    ws.append([1, "Concrete 240 ksc", "m3", 10, 20000, 5000, 25000])
    extractor = SheetExtractor(); classifier = RowClassifier(); builder = WBSBuilder()
    rows = [extractor.extract_row(ws, i, config()) for i in range(2, 5)]
    assert [classifier.classify(row) for row in rows] == [RowType.WBS, RowType.WBS, RowType.ITEM]
    builder.accept_heading(rows[0]); builder.accept_heading(rows[1])
    output = builder.build_item(rows[2], "BOQ")
    assert output[5:8] == ["Structural Work", "Concrete", "Concrete 240 ksc"]
