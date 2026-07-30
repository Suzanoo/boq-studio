from pathlib import Path

from openpyxl import Workbook, load_workbook

from boq_studio.core.exporter import normalize_workbook
from boq_studio.core.tree import DraftTree, DraftTreeNode, TreeNodeType
from boq_studio.core.tree_edit import DraftTreeEditor, TreeEditError
from boq_studio.models import NormalizeRequest, SheetConfig


def node(node_id, kind, description, *, row=None, amount=None, children=None):
    return DraftTreeNode(
        node_id=node_id,
        node_type=kind,
        description=description,
        source_row=row,
        amount=amount,
        children=children or [],
    )


def tree_with(*children):
    return DraftTree("ARCH", node("sheet:ARCH", TreeNodeType.SHEET_ROOT, "ARCH", children=list(children)))


def test_move_up_and_down_reorders_siblings():
    a = node("a", TreeNodeType.ITEM, "A")
    b = node("b", TreeNodeType.ITEM, "B")
    c = node("c", TreeNodeType.ITEM, "C")
    tree = tree_with(a, b, c)
    editor = DraftTreeEditor(tree)

    editor.move_up("b")
    assert [n.node_id for n in tree.root.children] == ["b", "a", "c"]

    editor.move_down("b")
    assert [n.node_id for n in tree.root.children] == ["a", "b", "c"]


def test_move_rejects_tree_edge():
    tree = tree_with(node("a", TreeNodeType.ITEM, "A"))
    editor = DraftTreeEditor(tree)
    try:
        editor.move_up("a")
    except TreeEditError as exc:
        assert "already first" in str(exc)
    else:
        raise AssertionError("Expected TreeEditError")


def test_delete_any_row_promotes_its_children():
    child = node("i", TreeNodeType.ITEM, "Concrete")
    source_wbs = node("w", TreeNodeType.WBS, "Floor", children=[child])
    tree = tree_with(source_wbs)

    DraftTreeEditor(tree).delete_node("w")

    assert tree.root.children == [child]


def test_tree_export_uses_order_deletion_and_wbs_depth(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "ARCH"
    ws.append(["No", "Description", "Unit", "Qty", "Amount"])
    ws.append(["1", "Floor", None, None, 300])
    ws.append(["1.1", "Item A", "m2", 1, 100])
    ws.append(["1.2", "Item B", "m2", 1, 200])
    wb.save(source)

    item_b = node("ARCH:4", TreeNodeType.ITEM, "Item B", row=4, amount=200)
    item_a = node("ARCH:3", TreeNodeType.ITEM, "Item A", row=3, amount=100)
    heading = node("ARCH:2", TreeNodeType.WBS, "Floor", row=2, amount=300, children=[item_b, item_a])
    tree = tree_with(heading)
    DraftTreeEditor(tree).delete_node("ARCH:3")

    cfg = SheetConfig(
        sheet_name="ARCH",
        header_row=1,
        start_row=2,
        item_no_col=1,
        description_cols=[2],
        unit_col=3,
        qty_col=4,
        amount_col=5,
    )
    request = NormalizeRequest(source, output, [cfg], draft_trees={"ARCH": tree})
    normalize_workbook(request)

    out = load_workbook(output, data_only=False)
    ws_out = out["ARCH"]
    assert ws_out.max_row == 2
    assert ws_out["F2"].value == "Floor"
    assert ws_out["H2"].value == "Item B"
    assert ws_out["M2"].value == 200
    out.close()
