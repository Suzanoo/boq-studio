from pathlib import Path

from openpyxl import Workbook, load_workbook

from boq_studio.core.exporter import OUTPUT_HEADERS, normalize_workbook
from boq_studio.core.tree import DraftTree, DraftTreeNode, TreeNodeType
from boq_studio.models import NormalizeRequest, SheetConfig


def _node(node_id: str, node_type: TreeNodeType, description: str, *, row=None, children=None):
    return DraftTreeNode(
        node_id=node_id,
        node_type=node_type,
        description=description,
        source_row=row,
        children=children or [],
    )


def _tree(sheet: str, heading: str, item_row: int) -> DraftTree:
    item = _node(f"{sheet}:{item_row}", TreeNodeType.ITEM, f"{sheet} Item", row=item_row)
    wbs = _node(f"{sheet}:wbs", TreeNodeType.WBS, heading, children=[item])
    root = _node(f"sheet:{sheet}", TreeNodeType.SHEET_ROOT, sheet, children=[wbs])
    return DraftTree(sheet, root)


def _config(sheet: str) -> SheetConfig:
    return SheetConfig(
        sheet_name=sheet,
        header_row=1,
        start_row=2,
        item_no_col=1,
        description_cols=[2],
        unit_col=3,
        qty_col=4,
        amount_col=5,
    )


def test_export_contract_contains_required_cost_columns():
    assert OUTPUT_HEADERS == [
        "WBS No.", "WBS-1", "WBS-2", "WBS-3", "WBS-4",
        "Description", "Unit", "Qty", "Material", "Labor", "Amount",
    ]


def test_tree_numbering_is_unique_across_sheets_and_project_uses_formulas(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    arch = wb.active
    arch.title = "ARCH"
    structure = wb.create_sheet("STRUCTURE")
    for ws, description, amount in (
        (arch, "ARCH Item", 100),
        (structure, "STRUCTURE Item", 200),
    ):
        ws.append(["No", "Description", "Unit", "Qty", "Amount"])
        ws.append([1, description, "lot", 1, amount])
    wb.save(source)

    request = NormalizeRequest(
        source,
        output,
        [_config("ARCH"), _config("STRUCTURE")],
        draft_trees={
            "ARCH": _tree("ARCH", "Architectural Work", 2),
            "STRUCTURE": _tree("STRUCTURE", "Structural Work", 2),
        },
    )
    normalize_workbook(request)

    out = load_workbook(output, data_only=False)
    assert [out["ARCH"][f"A{row}"].value for row in range(2, 5)] == ["1", "1.1", "1.1.1"]
    assert [out["STRUCTURE"][f"A{row}"].value for row in range(2, 5)] == ["2", "2.1", "2.1.1"]

    project = out["Project"]
    assert project["A2"].value == '=IF(\'ARCH\'!A2="","",\'ARCH\'!A2)'
    assert project["A5"].value == '=IF(\'STRUCTURE\'!A2="","",\'STRUCTURE\'!A2)'
    assert project["K4"].value == '=IF(\'ARCH\'!K4="","",\'ARCH\'!K4)'
    assert project["K7"].value == '=IF(\'STRUCTURE\'!K4="","",\'STRUCTURE\'!K4)'
    out.close()


def test_internal_node_id_never_appears_in_export(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ARCH"
    ws.append(["No", "Description", "Unit", "Qty", "Amount"])
    ws.append([1, "Door", "set", 1, 500])
    wb.save(source)

    tree = _tree("ARCH", "Door Work", 2)
    normalize_workbook(NormalizeRequest(source, output, [_config("ARCH")], {"ARCH": tree}))

    out = load_workbook(output, data_only=False)
    exported_values = {
        cell.value
        for row in out["ARCH"].iter_rows()
        for cell in row
        if cell.value not in (None, "")
    }
    assert "ARCH:2" not in exported_values
    assert "ARCH:wbs" not in exported_values
    out.close()
