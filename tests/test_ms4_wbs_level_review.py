from pathlib import Path

from openpyxl import Workbook

from boq_studio.core.classification import RowType
from boq_studio.core.exporter import normalize_sheet
from boq_studio.core.review import ClassificationReviewService
from boq_studio.models import SheetConfig
from boq_studio.templates import load_template, save_template


def make_sheet():
    wb = Workbook()
    ws = wb.active
    ws.append(["No", "Description", "Unit", "Qty", "Amount"])
    ws.append(["", "Building Zone", "", "", ""])
    ws.append(["", "Concrete Work", "", "", ""])
    ws.append(["", "Foundations", "", "", ""])
    ws.append(["", "Concrete item", "m3", 2, 100])
    return wb, ws


def config(**kwargs):
    values = dict(
        sheet_name="BOQ",
        header_row=1,
        start_row=2,
        item_no_col=1,
        description_cols=[2],
        unit_col=3,
        qty_col=4,
        amount_col=5,
    )
    values.update(kwargs)
    return SheetConfig(**values)


def test_review_exposes_effective_wbs_level():
    wb, ws = make_sheet()
    cfg = config(
        classification_overrides={2: RowType.WBS.value},
        wbs_level_overrides={2: 2},
    )
    row = ClassificationReviewService().build_rows(ws, cfg)[0]
    assert row.row_type == RowType.WBS
    assert row.wbs_level == 2
    assert row.level_overridden is True
    wb.close()


def test_export_applies_reviewed_wbs_2_3_4_hierarchy():
    wb, ws = make_sheet()
    out_wb = Workbook()
    out_ws = out_wb.active
    cfg = config(
        classification_overrides={
            2: RowType.WBS.value,
            3: RowType.WBS.value,
            4: RowType.WBS.value,
        },
        wbs_level_overrides={2: 2, 3: 3, 4: 4},
    )
    result = normalize_sheet(ws, out_ws, cfg)
    assert result.output_rows == 1
    assert out_ws["E2"].value == "Building Zone"
    assert out_ws["F2"].value == "Concrete Work"
    assert out_ws["G2"].value == "Foundations"
    assert out_ws["H2"].value == "Concrete item"
    assert out_ws["B2"].value == 4
    wb.close()
    out_wb.close()


def test_wbs_3_clears_previous_wbs_4():
    wb = Workbook()
    ws = wb.active
    ws.append(["No", "Description", "Unit", "Qty", "Amount"])
    ws.append(["", "Group A", "", "", ""])
    ws.append(["", "Sub A", "", "", ""])
    ws.append(["", "Item A", "m", 1, 10])
    ws.append(["", "Group B", "", "", ""])
    ws.append(["", "Item B", "m", 1, 20])
    out_wb = Workbook()
    out_ws = out_wb.active
    cfg = config(
        classification_overrides={2: "WBS", 3: "WBS", 5: "WBS"},
        wbs_level_overrides={2: 3, 3: 4, 5: 3},
    )
    normalize_sheet(ws, out_ws, cfg)
    assert out_ws["F2"].value == "Group A"
    assert out_ws["G2"].value == "Sub A"
    assert out_ws["F3"].value == "Group B"
    assert out_ws["G3"].value in (None, "")
    wb.close()
    out_wb.close()


def test_template_round_trip_preserves_wbs_levels(tmp_path: Path):
    path = tmp_path / "mapping.json"
    cfg = config(
        classification_overrides={2: "WBS", 3: "WBS"},
        wbs_level_overrides={2: 2, 3: 4},
    )
    save_template(path, cfg)
    loaded = load_template(path)
    assert loaded.wbs_level_overrides == {2: 2, 3: 4}
