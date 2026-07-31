from pathlib import Path

from openpyxl import Workbook, load_workbook

from boq_studio.core.exporter import normalize_workbook
from boq_studio.core.tree import DraftTree, DraftTreeNode, TreeNodeType
from boq_studio.models import NormalizeRequest, SheetConfig


def _node(node_id, node_type, description, row=None, children=None):
    return DraftTreeNode(node_id, node_type, description, source_row=row, children=children or [])


def test_export_adds_excel_outline_and_professional_page_setup(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BUILT-IN"
    ws.append(["No", "Description", "Unit", "Qty", "Amount"])
    ws.append([1, "Item A", "lot", 1, 100])
    ws.append([2, "Item B", "lot", 1, 200])
    wb.save(source)

    item_a = _node("item:a", TreeNodeType.ITEM, "Item A", 2)
    item_b = _node("item:b", TreeNodeType.ITEM, "Item B", 3)
    wbs4 = _node("wbs:4", TreeNodeType.WBS, "Wall Work", children=[item_a, item_b])
    wbs3 = _node("wbs:3", TreeNodeType.WBS, "Lobby", children=[wbs4])
    root = _node("root", TreeNodeType.SHEET_ROOT, "BUILT-IN", children=[wbs3])
    tree = DraftTree("BUILT-IN", root)
    cfg = SheetConfig(
        sheet_name="BUILT-IN", header_row=1, start_row=2,
        item_no_col=1, description_cols=[2], unit_col=3, qty_col=4, amount_col=5,
    )
    normalize_workbook(NormalizeRequest(source, output, [cfg], {"BUILT-IN": tree}))

    out = load_workbook(output, data_only=False)
    sheet = out["BUILT-IN"]
    assert [sheet.row_dimensions[r].outlineLevel for r in range(2, 7)] == [0, 1, 2, 3, 3]
    assert sheet.sheet_properties.outlinePr.summaryBelow is False
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:I6"
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.print_title_rows == "$1:$1"
    assert sheet.print_area == "'BUILT-IN'!$A$1:$I$6"
    assert sheet["A3"].font.bold is True
    assert sheet["A5"].font.bold is False

    project = out["Project"]
    assert [project.row_dimensions[r].outlineLevel for r in range(2, 7)] == [0, 1, 2, 3, 3]
    assert project.sheet_properties.outlinePr.summaryBelow is False
    out.close()
