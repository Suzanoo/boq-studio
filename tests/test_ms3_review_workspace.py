from pathlib import Path

from openpyxl import Workbook

from boq_studio.core.classification import RowType
from boq_studio.core.exporter import normalize_sheet
from boq_studio.core.review import ClassificationReviewService
from boq_studio.models import SheetConfig
from boq_studio.templates import load_template, save_template


def make_sheet():
    wb = Workbook()
    ws = wb.active
    ws.append(["No", "Description", "Unit", "Qty", "Amount"])
    ws.append(["", "Ambiguous heading", "", "", ""])
    ws.append(["", "Concrete item", "m3", 2, 100])
    return wb, ws


def config(**kwargs):
    values = dict(
        sheet_name="BOQ",
        header_row=1,
        start_row=2,
        item_no_col=1,
        description_cols=[2],
        unit_col=3,
        qty_col=4,
        amount_col=5,
    )
    values.update(kwargs)
    return SheetConfig(**values)


def test_review_rows_show_auto_and_effective_type():
    wb, ws = make_sheet()
    cfg = config(classification_overrides={2: RowType.IGNORE.value})
    rows = ClassificationReviewService().build_rows(ws, cfg)
    first = rows[0]
    assert first.auto_type == RowType.UNKNOWN
    assert first.row_type == RowType.IGNORE
    assert first.overridden is True
    wb.close()


def test_ignore_override_is_not_used_as_wbs_heading():
    wb, ws = make_sheet()
    output_wb = Workbook()
    output_ws = output_wb.active
    cfg = config(classification_overrides={2: RowType.IGNORE.value})
    result = normalize_sheet(ws, output_ws, cfg)
    assert result.output_rows == 1
    assert output_ws["F2"].value == "Concrete item"
    assert output_ws["C2"].value == "BOQ"
    assert output_ws["D2"].value in (None, "")
    wb.close()
    output_wb.close()


def test_unknown_can_be_confirmed_as_wbs():
    wb, ws = make_sheet()
    output_wb = Workbook()
    output_ws = output_wb.active
    cfg = config(classification_overrides={2: RowType.WBS.value})
    normalize_sheet(ws, output_ws, cfg)
    assert output_ws["D2"].value == "Ambiguous heading"
    wb.close()
    output_wb.close()


def test_template_round_trip_preserves_integer_override_rows(tmp_path: Path):
    path = tmp_path / "mapping.json"
    cfg = config(classification_overrides={2: RowType.WBS.value, 8: RowType.IGNORE.value})
    save_template(path, cfg)
    loaded = load_template(path)
    assert loaded.classification_overrides == {2: "WBS", 8: "IGNORE"}
