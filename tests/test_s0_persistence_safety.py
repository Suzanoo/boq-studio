"""S0 behavior regressions: real session files/XLSX, headless UI adapters."""
from __future__ import annotations

import copy
import json
import os
import queue
import subprocess
import sys
from pathlib import Path
from types import MethodType, SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from boq_studio import app, session, templates
from boq_studio.core.tree import DraftTree, DraftTreeNode, TreeNodeType
from boq_studio.core.tree_edit import DraftTreeEditor
from boq_studio.models import SheetConfig, SheetResult, NormalizeRequest
from boq_studio.core.exporter import normalize_workbook


class Var:
    def __init__(self, value=False): self.value = value
    def get(self): return self.value
    def set(self, value): self.value = value


class Widget:
    def __init__(self): self.options = {}; self.rows = {}; self.chosen = ('ARCH',)
    def configure(self, **kwargs): self.options.update(kwargs)
    def __setitem__(self, key, value): self.options[key] = value
    def selection(self): return self.chosen
    def get_children(self): return tuple(self.rows)
    def delete(self, *names):
        for name in names: self.rows.pop(name, None)
    def insert(self, _parent, _where, iid, values): self.rows[iid] = values
    def item(self, name, **kwargs): self.rows[name] = kwargs['values']


def make_workbook(path):
    wb = Workbook(); ws = wb.active; ws.title = 'ARCH'
    ws.append(['Item', 'Description', 'Unit', 'Qty', 'Amount'])
    ws.append(['1', 'Concrete', None, None, None])
    ws.append(['1.1', 'Slab', 'm3', 10, 1000])
    wb.save(path); wb.close()
    return path


def config():
    return SheetConfig('ARCH', item_no_col=1, description_cols=[2], unit_col=3,
                       qty_col=4, amount_col=5, classification_overrides={2: 'WBS'},
                       wbs_level_overrides={2: 3})


def tree():
    return DraftTree('ARCH', DraftTreeNode('sheet:ARCH', TreeNodeType.SHEET_ROOT, 'ARCH', children=[
        DraftTreeNode('virtual:1', TreeNodeType.VIRTUAL_WBS, 'Reviewed งาน', children=[
            DraftTreeNode('ARCH:3', TreeNodeType.ITEM, 'Reviewed slab', source_row=3, amount=1000)
        ])
    ]))


def payload(source):
    return session.build_session_payload(source, {'ARCH': True}, {'ARCH': config()},
                                         {'ARCH': tree()}, {'ARCH'}, set())


@pytest.fixture
def ui(tmp_path, monkeypatch):
    monkeypatch.setattr(session, 'session_directory', lambda: tmp_path / 'sessions')
    monkeypatch.setattr(templates, 'template_directory', lambda: tmp_path / 'templates')
    monkeypatch.setattr(app.tk, 'BooleanVar', Var)
    messages = []
    for name in ('showerror', 'showwarning', 'showinfo'):
        monkeypatch.setattr(app.messagebox, name, lambda *a, **kw: messages.append(a))
    monkeypatch.setattr(app.messagebox, 'askyesno', lambda *a, **kw: True)
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: True)
    monkeypatch.setattr(app.filedialog, "asksaveasfilename", lambda **kw: "")
    h = SimpleNamespace(input_path=None, sheet_vars={}, configs={}, draft_trees={},
                        tree_ready_sheets=set(), exported_sheets=set(), is_running=False,
                        _session_dirty=False, _session_token=None, _source_info={},
                        events=queue.Queue(), messages=messages, closed=False)
    for name, value in vars(app.BOQStudioApp).items():
        if name.startswith('_') and not name.startswith('__') and callable(value):
            setattr(h, name, MethodType(value, h))
    for name in ('configure_selected', 'open_draft_tree', 'run_normalize', 'toggle_sheet'):
        setattr(h, name, MethodType(getattr(app.BOQStudioApp, name), h))
    for name in ('file_label', 'sheet_tree', 'select_all_button', 'clear_button',
                 'configure_button', 'tree_button', 'run_button', 'status_label', 'progress'):
        setattr(h, name, Widget())
    h._remember_recent_file = lambda path: None
    h.wait_window = lambda dialog: None
    h.after = lambda *a: None
    h.destroy = lambda: setattr(h, 'closed', True)
    h._load_workbook_path(make_workbook(tmp_path / 'source.xlsx'))
    h.configs = {'ARCH': config()}; h.draft_trees = {'ARCH': tree()}
    h.tree_ready_sheets = {'ARCH'}
    assert h._save_project_session()
    return h


def disk_path(h): return session.session_path_for(h.input_path)


def test_version_one_roundtrip_and_backup(tmp_path):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    first = payload(source); token = session.save_session(path, first, expected_token=None)
    second = copy.deepcopy(first); second['trees']['ARCH']['root']['description'] = 'new'
    session.save_session(path, second, expected_token=token)
    assert session.load_session(path) == json.loads(json.dumps(second))
    assert session.load_session(path.with_suffix('.json.bak')) == json.loads(json.dumps(first))
    assert session.config_from_dict(first['configs']['ARCH']) == config()
    assert session.tree_to_dict(session.tree_from_dict(first['trees']['ARCH'])) == first['trees']['ARCH']


def test_noop_save_does_not_rotate_backup(tmp_path):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); token = session.save_session(path, p)
    p['saved_at'] = 'different timestamp'
    assert session.save_session(path, p, expected_token=token) == token
    assert not path.with_suffix('.json.bak').exists()


@pytest.mark.parametrize('point', ['fsync', 'backup', 'primary'])
def test_failed_save_keeps_last_good_primary(tmp_path, monkeypatch, point):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); token = session.save_session(path, p); old = path.read_bytes()
    p['selected']['ARCH'] = False
    original = Path.replace
    def fail_replace(self, target):
        if (point == 'backup' and str(target).endswith('.bak')) or (point == 'primary' and target == path):
            raise OSError('simulated locked/full disk')
        return original(self, target)
    if point == 'fsync':
        monkeypatch.setattr(session.os, 'fsync', lambda fd: (_ for _ in ()).throw(OSError('flush failed')))
    else:
        monkeypatch.setattr(Path, 'replace', fail_replace)
    with pytest.raises(OSError): session.save_session(path, p, expected_token=token)
    assert path.read_bytes() == old
    assert not list(tmp_path.glob('*.tmp'))


def test_corrupt_primary_does_not_destroy_valid_backup(tmp_path):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); session.save_session(path, p)
    p['selected']['ARCH'] = False; session.save_session(path, p)
    backup = path.with_suffix('.json.bak').read_bytes()
    path.write_text('broken')
    session.save_session(path, p, expected_token=session.session_token(path))
    assert path.with_suffix('.json.bak').read_bytes() == backup
    assert any(f.read_text() == 'broken' for f in tmp_path.glob('*.recovery-*.json'))


def test_stale_writer_cannot_overwrite_or_rotate_backup(tmp_path):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); stale = session.save_session(path, p)
    newer = copy.deepcopy(p); newer['selected']['ARCH'] = False
    session.save_session(path, newer, expected_token=stale)
    old = path.read_bytes(); backup = path.with_suffix('.json.bak').read_bytes()
    with pytest.raises(session.SessionConflictError): session.save_session(path, p, expected_token=stale)
    assert path.read_bytes() == old and path.with_suffix('.json.bak').read_bytes() == backup


@pytest.mark.parametrize('mutation', [
    lambda p: p.update(version=999),
    lambda p: p['tree_ready'].append('MISSING'),
    lambda p: p['trees']['ARCH']['root']['children'].append(copy.deepcopy(p['trees']['ARCH']['root']['children'][0])),
    lambda p: p['configs']['ARCH'].update(description_cols=['bad']),
    lambda p: p['trees']['ARCH']['root'].update(amount=float('nan')),
    lambda p: p.update(tree_ready=[], exported=['ARCH']),
])
def test_invalid_payload_never_replaces_saved_work(tmp_path, mutation):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); session.save_session(path, p); old = path.read_bytes()
    mutation(p)
    with pytest.raises((ValueError, TypeError)): session.save_session(path, p)
    assert path.read_bytes() == old


def test_archive_survives_later_saves(tmp_path):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); token = session.save_session(path, p)
    archived = session.archive_session(path, expected_token=token); old = archived.read_bytes()
    for value in (False, True, False):
        p['selected']['ARCH'] = value; session.save_session(path, p)
    assert archived.read_bytes() == old


def test_process_interruption_before_commit_retains_primary(tmp_path):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); session.save_session(path, p); old = path.read_bytes()
    code = '''
import os,sys
from pathlib import Path
from boq_studio.session import load_session,save_session
path=Path(sys.argv[1]); p=load_session(path); p['selected']['ARCH']=False
replace=Path.replace
def die(self,target):
    if target == path: os._exit(17)
    return replace(self,target)
Path.replace=die
save_session(path,p)
'''
    env = dict(os.environ, PYTHONPATH=str(Path(__file__).resolve().parents[1] / 'src'))
    result = subprocess.run([sys.executable, '-c', code, str(path)], env=env, check=False)
    assert result.returncode == 17 and path.read_bytes() == old
    assert session.load_session(path.with_suffix('.json.bak')) == json.loads(json.dumps(p))
    # Process exit releases the OS lock, even though the .lock filename remains.
    p['selected']['ARCH'] = False; session.save_session(path, p)
    assert session.load_session(path)['selected']['ARCH'] is False


def test_resume_restores_exact_reviewed_work(ui):
    expected = session.tree_to_dict(ui.draft_trees['ARCH'])
    ui._load_workbook_path(ui.input_path)
    assert session.tree_to_dict(ui.draft_trees['ARCH']) == expected
    assert ui.tree_ready_sheets == {'ARCH'}
    assert ui.configs['ARCH'] == config()


def test_failed_open_preserves_path_and_state(ui, tmp_path):
    old_path = ui.input_path; old_tree = ui.draft_trees
    invalid = tmp_path / 'broken.xlsx'; invalid.write_text('bad')
    ui._load_workbook_path(invalid)
    assert ui.input_path == old_path and ui.draft_trees is old_tree


def test_resume_cancel_keeps_current_project(ui, tmp_path, monkeypatch):
    other = make_workbook(tmp_path / 'other.xlsx')
    session.save_session(session.session_path_for(other), payload(other))
    old_path = ui.input_path; old_tree = ui.draft_trees
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: None)
    ui._load_workbook_path(other)
    assert ui.input_path == old_path and ui.draft_trees is old_tree


@pytest.mark.parametrize('kind', ['decline', 'changed', 'corrupt', 'newer'])
def test_start_fresh_preserves_old_bytes_before_overwrite(ui, monkeypatch, kind):
    path = disk_path(ui)
    if kind == 'changed':
        wb = load_workbook(ui.input_path); wb['ARCH']['D3'] = 99; wb.save(ui.input_path); wb.close()
    elif kind == 'corrupt': path.write_text('broken')
    elif kind == 'newer':
        p = json.loads(path.read_text()); p['version'] = 999; path.write_text(json.dumps(p))
    old = path.read_bytes()
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: False)
    ui._load_workbook_path(ui.input_path)
    assert ui.draft_trees == {} and ui.tree_ready_sheets == set()
    assert ui._save_project_session()
    assert any(p.read_bytes() == old for p in path.parent.glob('*.recovery-*.json'))


def test_archive_failure_cancels_fresh_open(ui, monkeypatch):
    old = ui.draft_trees; old_bytes = disk_path(ui).read_bytes()
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: False)
    monkeypatch.setattr(app, 'archive_session', lambda *a, **kw: (_ for _ in ()).throw(OSError('full')))
    ui._load_workbook_path(ui.input_path)
    assert ui.draft_trees is old and disk_path(ui).read_bytes() == old_bytes


def test_corrupt_session_can_resume_backup(ui):
    path = disk_path(ui); original = path.read_bytes()
    path.with_suffix('.json.bak').write_bytes(original); path.write_text('bad json')
    ui._load_workbook_path(ui.input_path)
    assert ui.tree_ready_sheets == {'ARCH'} and ui._session_dirty
    ui._request_close()
    assert ui.closed and session.load_session(path)['trees']['ARCH']['root']['children'][0]['description'] == 'Reviewed งาน'


def test_unrelated_template_cannot_restore_row_overrides(ui, tmp_path, monkeypatch):
    other = make_workbook(tmp_path / 'other.xlsx')
    templates.save_template(templates.default_template_path(other, 'ARCH'), config())
    ui._load_workbook_path(other)
    assert ui.configs['ARCH'].description_cols == [2]
    assert ui.configs['ARCH'].classification_overrides == {}
    assert ui.configs['ARCH'].wbs_level_overrides == {}


def test_mapping_unchanged_preserves_tree_approval_and_session(ui, monkeypatch):
    tree_before = ui.draft_trees['ARCH']; old = disk_path(ui).read_bytes()
    monkeypatch.setattr(app, 'SheetMappingDialog', lambda *a: SimpleNamespace(result=copy.deepcopy(ui.configs['ARCH'])))
    ui.configure_selected()
    assert ui.draft_trees['ARCH'] is tree_before and ui.tree_ready_sheets == {'ARCH'}
    assert disk_path(ui).read_bytes() == old


@pytest.mark.parametrize('accept', [True, False])
def test_mapping_change_requires_consent_and_archives_reviewed_tree(ui, monkeypatch, accept):
    old = disk_path(ui).read_bytes(); cfg = copy.deepcopy(ui.configs['ARCH']); cfg.start_row = 3
    monkeypatch.setattr(app, 'SheetMappingDialog', lambda *a: SimpleNamespace(result=cfg))
    monkeypatch.setattr(app.messagebox, 'askyesno', lambda *a, **kw: accept)
    ui.configure_selected()
    if accept:
        assert ui.draft_trees == {} and ui.tree_ready_sheets == set() and ui.exported_sheets == set()
        assert any(p.read_bytes() == old for p in disk_path(ui).parent.glob('*.recovery-*.json'))
    else:
        assert disk_path(ui).read_bytes() == old and ui.tree_ready_sheets == {'ARCH'}
        assert ui.configs['ARCH'].start_row == 2


def test_failed_tree_commit_keeps_parent_and_editor_candidate(ui, monkeypatch):
    parent = ui.draft_trees['ARCH']; candidate = copy.deepcopy(parent)
    candidate.root.children[0].description = 'Unsaved candidate'
    monkeypatch.setattr(app, 'save_session', lambda *a, **kw: (_ for _ in ()).throw(OSError('disk full')))
    assert not ui._commit_tree('ARCH', candidate, 'draft')
    assert ui.draft_trees['ARCH'] is parent and ui.tree_ready_sheets == {'ARCH'}
    assert candidate.root.children[0].description == 'Unsaved candidate'


def test_tree_save_draft_and_approval_survive_resume(ui):
    candidate = copy.deepcopy(ui.draft_trees['ARCH']); candidate.root.children.reverse()
    assert ui._commit_tree('ARCH', candidate, 'draft')
    ui._load_workbook_path(ui.input_path)
    assert ui.tree_ready_sheets == set() and ui.draft_trees
    assert ui._commit_tree('ARCH', ui.draft_trees['ARCH'], 'ready')
    ui._load_workbook_path(ui.input_path)
    assert ui.tree_ready_sheets == {'ARCH'}


def test_real_dialog_initialization_isolates_discarded_edits(monkeypatch):
    class StopInit(Exception): pass
    captured = []
    monkeypatch.setattr(app.WorkspaceWindow, '__init__', lambda *a, **kw: None)
    monkeypatch.setattr(app, 'LayoutManager', lambda *a, **kw: None)
    def stop(tree, **kw): captured.append(tree); raise StopInit()
    monkeypatch.setattr(app, 'DraftTreeEditor', stop)
    original = tree(); dialog = object.__new__(app.DraftTreeDialog)
    with pytest.raises(StopInit):
        app.DraftTreeDialog.__init__(dialog, None, Path('source.xlsx'), config(), original)
    assert captured[0] is not original
    DraftTreeEditor(dialog.draft_tree).rename_node('virtual:1', 'Discard me')
    dialog.dirty = True; dialog._persist_view_state = lambda: None; dialog.destroy = lambda: None
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: False)
    dialog._cancel()
    assert original.root.children[0].description == 'Reviewed งาน' and dialog.result is None


def test_failed_editor_save_does_not_close_or_approve():
    dialog = SimpleNamespace(on_commit=lambda tree, mode: False, draft_tree=tree(),
                             dirty=True, result=None, closed=False)
    dialog.destroy = lambda: setattr(dialog, 'closed', True)
    app.DraftTreeDialog._finish_tree(dialog, 'ready')
    assert not dialog.closed and dialog.dirty and dialog.result is None


def test_source_change_blocks_save_and_does_not_restamp_hash(ui):
    path = disk_path(ui); old = path.read_bytes()
    ui.input_path.write_bytes(b'changed source')
    ui._session_dirty = True
    assert not ui._save_project_session()
    assert path.read_bytes() == old
    assert ui._session_payload()['source']['sha256'] == json.loads(old)['source']['sha256']


def test_failed_save_blocks_close_and_switch(ui, tmp_path, monkeypatch):
    old_path = ui.input_path; ui._session_dirty = True
    monkeypatch.setattr(app, 'save_session', lambda *a, **kw: (_ for _ in ()).throw(OSError('full')))
    monkeypatch.setattr(app.messagebox, 'askyesno', lambda *a, **kw: False)
    ui._request_close()
    ui._load_workbook_path(make_workbook(tmp_path / 'other.xlsx'))
    assert not ui.closed and ui.input_path == old_path and ui._session_dirty


def test_source_change_can_save_separate_recovery_before_close(ui, tmp_path, monkeypatch):
    old = disk_path(ui).read_bytes(); ui.input_path.write_bytes(b'changed')
    ui.draft_trees['ARCH'].root.children[0].description = 'Latest recovered work'
    ui._session_dirty = True
    target = tmp_path / 'manual-recovery.json'
    monkeypatch.setattr(app.filedialog, 'asksaveasfilename', lambda **kw: str(target))
    ui._request_close()
    assert ui.closed and disk_path(ui).read_bytes() == old
    p = session.load_session(target)
    assert p['source']['sha256'] == json.loads(old)['source']['sha256']
    assert p['trees']['ARCH']['root']['children'][0]['description'] == 'Latest recovered work'


def test_export_requires_actual_tree_not_stale_ready_flag(ui, monkeypatch):
    ui.draft_trees.clear()
    monkeypatch.setattr(app.threading, 'Thread', lambda **kw: pytest.fail('must not start export'))
    ui.run_normalize()
    assert any('Draft WBS Tree required' in m for m in ui.messages)


def test_busy_guards_block_mutating_commands_and_close(ui, tmp_path):
    ui.is_running = True; original = ui.input_path
    ui._load_workbook_path(tmp_path / 'other.xlsx'); ui._set_all(False)
    ui.configure_selected(); ui.open_draft_tree(); ui._request_close(); ui.run_normalize()
    assert ui.input_path == original and ui.sheet_vars['ARCH'].get() and not ui.closed


def test_export_completion_records_actual_success_only(ui):
    ui.events.put(('done', Path('output.xlsx'), [SheetResult('ARCH', 'ERROR')], ('ARCH',)))
    ui.exported_sheets = {'ARCH'}; ui.is_running = True
    ui._poll_events()
    assert not ui.is_running and ui.exported_sheets == set()
    assert session.load_session(disk_path(ui))['exported'] == []


def test_snapshot_export_preserves_original_provenance_and_values(tmp_path):
    original = make_workbook(tmp_path / 'original.xlsx'); snapshot = tmp_path / 'copy.xlsx'
    snapshot.write_bytes(original.read_bytes())
    original.write_bytes(b'externally replaced while export runs')
    output = tmp_path / 'output.xlsx'
    results = normalize_workbook(NormalizeRequest(original, output, [config()], {'ARCH': tree()}), source_snapshot=snapshot)
    assert results[0].status == 'OK'
    wb = load_workbook(output)
    assert wb['LOG']['B2'].value == str(original)
    assert wb['ARCH']['K4'].value == 1000
    wb.close()


def test_failed_tree_candidate_can_be_recovered_without_changing_parent(ui, tmp_path, monkeypatch):
    original = copy.deepcopy(ui.draft_trees['ARCH'])
    candidate = copy.deepcopy(original); candidate.root.children[0].description = 'Latest candidate'
    ui.input_path.write_bytes(b'changed externally')
    target = tmp_path / 'candidate-recovery.json'
    monkeypatch.setattr(app.filedialog, 'asksaveasfilename', lambda **kw: str(target))
    assert not ui._commit_tree('ARCH', candidate, 'draft')
    assert session.tree_to_dict(ui.draft_trees['ARCH']) == session.tree_to_dict(original)
    recovered = session.load_session(target)
    assert recovered['trees']['ARCH']['root']['children'][0]['description'] == 'Latest candidate'
    assert recovered['tree_ready'] == []


def test_session_lock_rejects_second_writer(tmp_path):
    source = make_workbook(tmp_path / 'source.xlsx'); path = tmp_path / 'session.json'
    p = payload(source); session.save_session(path, p); old = path.read_bytes()
    code = '''
import sys
from pathlib import Path
from boq_studio.session import load_session,save_session,SessionConflictError
p=Path(sys.argv[1])
try: save_session(p,load_session(p))
except SessionConflictError: sys.exit(23)
sys.exit(0)
'''
    env = dict(os.environ, PYTHONPATH=str(Path(__file__).resolve().parents[1] / 'src'))
    with session._session_lock(path):
        result = subprocess.run([sys.executable, '-c', code, str(path)], env=env, check=False)
    assert result.returncode == 23 and path.read_bytes() == old


def test_export_worker_uses_verified_snapshot_and_original_log(ui, tmp_path, monkeypatch):
    monkeypatch.setattr(Path, 'home', lambda: tmp_path)
    real_export = app.normalize_workbook
    def export_then_replace_source(request, progress, *, source_snapshot):
        # Simulate Excel editing after worker verified its private copy.
        request.input_path.write_bytes(b'changed after snapshot validation')
        return real_export(request, progress, source_snapshot=source_snapshot)
    monkeypatch.setattr(app, 'normalize_workbook', export_then_replace_source)
    monkeypatch.setattr(app.threading, 'Thread', lambda target, **kw: SimpleNamespace(start=target))
    ui.run_normalize()
    assert ui.is_running
    ui._poll_events()
    output = tmp_path / 'Desktop' / 'source_OKD.xlsx'
    wb = load_workbook(output)
    assert wb['ARCH']['K4'].value == 1000 and wb['LOG']['B2'].value == str(ui.input_path)
    wb.close()
    assert ui._session_dirty and 'NOT saved' in ui.status_label.options['text']
    assert session.load_session(disk_path(ui))['exported'] == []


def test_export_rejects_change_before_snapshot_copy(ui, tmp_path, monkeypatch):
    monkeypatch.setattr(Path, 'home', lambda: tmp_path)
    real_copy = app.shutil.copyfile
    def copy_changed(src, dst):
        src.write_bytes(b'changed during export start')
        return real_copy(src, dst)
    monkeypatch.setattr(app.shutil, 'copyfile', copy_changed)
    monkeypatch.setattr(app.threading, 'Thread', lambda target, **kw: SimpleNamespace(start=target))
    ui.run_normalize(); ui._poll_events()
    assert not (tmp_path / 'Desktop' / 'source_OKD.xlsx').exists()
    assert any('Normalization failed' in m for m in ui.messages)


def test_failed_recovery_write_cannot_close(ui, tmp_path, monkeypatch):
    ui._session_dirty = True
    monkeypatch.setattr(app, 'save_session', lambda *a, **kw: (_ for _ in ()).throw(OSError('full')))
    monkeypatch.setattr(app.filedialog, 'asksaveasfilename', lambda **kw: str(tmp_path / 'recovery.json'))
    ui._request_close()
    assert not ui.closed and ui._session_dirty


def test_save_warning_result_remains_ready_not_exported(ui):
    ui.events.put(('done', Path('output.xlsx'), [SheetResult('ARCH', 'WARNING')], ('ARCH',)))
    ui.is_running = True; ui._poll_events()
    assert ui.exported_sheets == set() and ui.tree_ready_sheets == {'ARCH'}


def test_save_good_export_state_persists(ui):
    ui.events.put(('done', Path('output.xlsx'), [SheetResult('ARCH', 'OK')], ('ARCH',)))
    ui.is_running = True; ui._poll_events()
    assert session.load_session(disk_path(ui))['exported'] == ['ARCH']


def test_cancel_after_undo_redo_preserves_parent(ui, monkeypatch):
    # Existing dialog initialization isolation is tested above; exercise its real history methods.
    parent = ui.draft_trees['ARCH']
    dialog = SimpleNamespace(draft_tree=copy.deepcopy(parent), undo_stack=[], redo_stack=[],
                             dirty=False, result=None, _render=lambda *a, **kw: None,
                             _persist_view_state=lambda: None, destroy=lambda: None)
    for name in ('_remember', '_restore_tree', '_undo', '_redo', '_cancel'):
        setattr(dialog, name, MethodType(getattr(app.DraftTreeDialog, name), dialog))
    dialog._remember()
    DraftTreeEditor(dialog.draft_tree).rename_node('virtual:1', 'changed')
    dialog.dirty = True; dialog._undo(); dialog._redo()
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: False)
    dialog._cancel()
    assert parent.root.children[0].description == 'Reviewed งาน'


@pytest.mark.parametrize('operation', ['_save_draft', '_cancel'])
def test_view_preference_failure_does_not_block_save_or_discard(monkeypatch, operation):
    messages = []
    monkeypatch.setattr(app.messagebox, 'showwarning', lambda title, *a, **kw: messages.append(title))
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: False)
    def fail(*a, **kw): raise OSError('preferences read-only')
    dialog = SimpleNamespace(tree=SimpleNamespace(winfo_exists=lambda: True),
                             view_state=SimpleNamespace(capture=fail), geometry=lambda: '900x600',
                             draft_tree=object(), on_commit=lambda *a: True, dirty=True,
                             result=None, closed=False)
    dialog.destroy = lambda: setattr(dialog, 'closed', True)
    for name in ('_persist_view_state', '_finish_tree', '_save_draft', '_cancel'):
        setattr(dialog, name, MethodType(getattr(app.DraftTreeDialog, name), dialog))
    getattr(dialog, operation)()
    assert dialog.closed and messages == ['View preferences not saved']
    assert (dialog.result is dialog.draft_tree) == (operation == '_save_draft')
